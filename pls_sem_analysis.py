#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Complete PLS-SEM analysis pipeline
==================================

Financial control (Control Activities, Financial Management) and four
dimensions of public accountability (Organizational, Legal, Professional,
Political) - Tax Affairs Administration of Tehran.

Reads ``1.xlsx`` from the Desktop (or a path given on the command line) and
writes every table, figure and diagnostic to ``Desktop/PLS_SEM_Results``.

What is computed
----------------
Data screening      missing values, out-of-range codes, straight-lining,
                    Mahalanobis outliers
Descriptives        demographic profile, item and construct statistics
Normality           skewness, kurtosis, Shapiro-Wilk, Kolmogorov-Smirnov,
                    Mardia multivariate skewness / kurtosis, Q-Q plots
Dimensionality      KMO, Bartlett, eigenvalues, parallel analysis, EFA
                    (principal axis, varimax + promax), Harman single factor
Measurement model   outer loadings / weights (+ bootstrap), indicator
                    reliability, Cronbach's alpha, rho_A, rho_C, AVE,
                    outer VIF, cross-loadings, Fornell-Larcker, HTMT, HTMT2,
                    bootstrap HTMT inference, item-level diagnostics
Structural model    path coefficients, bootstrap (percentile, BC, BCa CIs),
                    t / p values, total and indirect effects, R2, adjusted R2,
                    f2, q2, inner VIF, Q2 (blindfolding), PLSpredict
                    (Q2_predict, RMSE, MAE vs. LM benchmark), post-hoc power,
                    minimum sample size (inverse square root method)
Model fit           SRMR, d_ULS, d_G, Chi-square, NFI, RMS_theta, GoF for
                    saturated and estimated models, Bollen-Stine bootstrap
                    exact-fit tests (HI95 / HI99); optional CB-SEM CFA fit
                    (CFI, TLI, RMSEA, ...) through semopy
CB-SEM (semopy)     CFA and structural model by maximum likelihood; Table 8
                    (X2/df, GFI, SRMR, IFI, NFI, PGFI, PNFI, RMSEA, CFI) plus
                    AGFI, RFI, TLI, PCFI, RMSEA 90% CI, PCLOSE, AIC, BIC;
                    CFA reliability; ML/ULS/GLS estimator comparison; PLS vs
                    CB-SEM triangulation; semopy HTML report
3D figures          response surfaces, 3D path map, 3D HTMT, 3D LV scatter,
                    3D cross-loadings; interactive rotatable HTML (plotly)
Common method bias  Harman single factor, full collinearity VIF (Kock 2015)
Robustness          Gaussian copula endogeneity test, nonlinear (quadratic)
                    effects, multigroup analysis with permutation test and
                    MICOM, case-influence (jackknife), outlier-excluded model
Sensitivity         consistent PLS (PLSc), alternative inner weighting
                    schemes, bootstrap seed stability, sum-score OLS (HC3),
                    demographic control variables, mediation model
                    (CA -> FM -> accountability), higher-order models,
                    item-purified model (loadings + HTMT driven)
IPMA                construct- and indicator-level importance-performance maps

Requirements
------------
    pip install numpy pandas scipy matplotlib openpyxl
    pip install semopy        (CB-SEM, Table 8)
Optional:
    pip install python-docx   (Word report)
    pip install plotly        (interactive 3D HTML)

Usage
-----
    python pls_sem_analysis.py                 # uses Desktop/1.xlsx
    python pls_sem_analysis.py "D:/data/1.xlsx"
"""

from __future__ import annotations

import math
import os
import sys
import time
import warnings
from collections import OrderedDict
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np
import pandas as pd
from scipy import stats

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch, Ellipse  # noqa: E402
from mpl_toolkits.mplot3d import Axes3D  # noqa: E402,F401

warnings.filterwarnings("ignore", category=RuntimeWarning)

# =============================================================================
# 1. CONFIGURATION - edit here only
# =============================================================================

INPUT_FILE_NAME = "1.xlsx"
SHEET_NAME = 0                      # first sheet
OUTPUT_FOLDER_NAME = "PLS_SEM_Results"

CONSTRUCTS = OrderedDict([
    ("CA",  {"name": "Control Activities",          "items": [f"Q{i}" for i in range(1, 10)]}),
    ("FM",  {"name": "Financial Management",        "items": [f"Q{i}" for i in range(10, 14)]}),
    ("ORG", {"name": "Organizational Accountability", "items": [f"Z{i}" for i in range(1, 8)]}),
    ("LEG", {"name": "Legal Accountability",        "items": [f"Z{i}" for i in range(8, 14)]}),
    ("PRO", {"name": "Professional Accountability", "items": [f"Z{i}" for i in range(14, 21)]}),
    ("POL", {"name": "Political Accountability",    "items": [f"Z{i}" for i in range(21, 25)]}),
])

# Hypothesised structural paths (source, target, label)
HYPOTHESES = [
    ("CA", "ORG", "H1"), ("CA", "LEG", "H2"), ("CA", "PRO", "H3"), ("CA", "POL", "H4"),
    ("FM", "ORG", "H5"), ("FM", "LEG", "H6"), ("FM", "PRO", "H7"), ("FM", "POL", "H8"),
]

# Items to exclude from the MAIN model (fill after reviewing the purification
# suggestions, e.g. ["Z9", "Z16"]). Leave empty for the full instrument.
DROP_ITEMS: list[str] = []

DEMOGRAPHICS = OrderedDict([
    ("Gender",     {1: "Male", 2: "Female"}),
    ("Age",        {1: "20-25 years", 2: "26-30 years", 3: "31-35 years", 4: "Above 35 years"}),
    ("Degree",     {1: "High school diploma or lower", 2: "Associate's degree",
                    3: "Bachelor's degree", 4: "Master's degree or higher"}),
    ("Experience", {1: "Less than 5 years", 2: "6 to 10 years", 3: "11 to 15 years",
                    4: "More than 15 years"}),
])

# Multigroup analysis: variable -> {group label: [codes]}
MGA_GROUPS = OrderedDict([
    ("Gender",     OrderedDict([("Male", [1]), ("Female", [2])])),
    ("Experience", OrderedDict([("10 years or less", [1, 2]), ("More than 10 years", [3, 4])])),
])
MGA_MIN_GROUP_SIZE = 30

SCALE_MIN, SCALE_MAX = 1, 5          # Likert range
MISSING_STRATEGY = "mean"            # "mean" (SmartPLS mean replacement) or "listwise"

INNER_WEIGHTING = "path"             # "path" (SmartPLS default), "factorial", "centroid"
MAX_ITER = 3000
TOLERANCE = 1e-7

N_BOOT = 5000                        # main bootstrap
N_BOOT_SENS = 2000                   # bootstraps for sensitivity models
N_BOOT_FIT = 1000                    # Bollen-Stine bootstrap for exact fit
N_BOOT_ROBUST = 1000                 # copula / nonlinear OLS bootstraps
N_PERM = 1000                        # MGA / MICOM permutations
N_PARALLEL = 500                     # parallel analysis replications
BLINDFOLD_D = 7
PREDICT_FOLDS, PREDICT_REPEATS = 10, 10
SEED = 2024
ALPHA = 0.05

# Purification (sensitivity) rules
PURIFY_MIN_ITEMS = 3
PURIFY_HTMT_TARGET = 0.90
PURIFY_MAX_SHARE = 0.40              # never drop more than 40% of a construct

FAST_MODE = os.environ.get("PLS_FAST", "0") == "1"   # True -> fewer resamples (quick test run)

DPI = 300

# Colour roles (validated reference categorical palette, light mode)
C_SERIES = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#008300", "#4a3aa7", "#e34948"]
C_TEXT, C_TEXT2, C_GRID, C_SURF = "#0b0b0b", "#52514e", "#dddcd7", "#fcfcfb"
C_BLUE_SEQ = ["#cde2fb", "#9ec5f4", "#6da7ec", "#3987e5", "#256abf", "#184f95", "#0d366b"]
C_DIV_NEG, C_DIV_MID, C_DIV_POS = "#e34948", "#f0efec", "#2a78d6"

if FAST_MODE:
    N_BOOT, N_BOOT_SENS, N_BOOT_FIT, N_BOOT_ROBUST, N_PERM, N_PARALLEL = 300, 200, 100, 200, 100, 100
    PREDICT_REPEATS = 2


# =============================================================================
# 2. GENERAL UTILITIES
# =============================================================================

def log(msg: str) -> None:
    print(f"[{time.strftime('%H:%M:%S')}] {msg}", flush=True)


def stars(p: float) -> str:
    if p is None or not np.isfinite(p):
        return ""
    return "***" if p < 0.001 else "**" if p < 0.01 else "*" if p < 0.05 else "n.s."


def fmt_p(p: float) -> str:
    if p is None or not np.isfinite(p):
        return "-"
    return "< 0.001" if p < 0.001 else f"{p:.3f}"


def standardize(X: np.ndarray, ddof: int = 0):
    m = X.mean(axis=0)
    s = X.std(axis=0, ddof=ddof)
    s = np.where(s == 0, 1.0, s)
    return (X - m) / s, m, s


def corr(X: np.ndarray) -> np.ndarray:
    Z, _, _ = standardize(X)
    return (Z.T @ Z) / Z.shape[0]


def nearest_pd(A: np.ndarray, eps: float = 1e-8) -> np.ndarray:
    A = (A + A.T) / 2
    vals, vecs = np.linalg.eigh(A)
    vals = np.clip(vals, eps, None)
    B = vecs @ np.diag(vals) @ vecs.T
    d = np.sqrt(np.diag(B))
    return B / np.outer(d, d)


def mat_power(A: np.ndarray, power: float) -> np.ndarray:
    vals, vecs = np.linalg.eigh((A + A.T) / 2)
    vals = np.clip(vals, 1e-10, None)
    return vecs @ np.diag(vals ** power) @ vecs.T


def p_from_t(t: float, df: int) -> float:
    return float(2 * stats.t.sf(abs(t), df)) if np.isfinite(t) else np.nan


def ols(y: np.ndarray, X: np.ndarray, add_const: bool = True):
    """OLS with HC3 robust standard errors."""
    if add_const:
        X = np.column_stack([np.ones(len(y)), X])
    XtX_inv = np.linalg.pinv(X.T @ X)
    b = XtX_inv @ X.T @ y
    e = y - X @ b
    h = np.einsum("ij,jk,ik->i", X, XtX_inv, X)
    w = (e / np.clip(1 - h, 1e-8, None)) ** 2
    cov = XtX_inv @ (X.T * w) @ X @ XtX_inv
    se = np.sqrt(np.diag(cov))
    r2 = 1 - (e @ e) / ((y - y.mean()) @ (y - y.mean()))
    return b, se, r2


def effect_label_f2(f2: float) -> str:
    if not np.isfinite(f2):
        return "-"
    return "large" if f2 >= 0.35 else "medium" if f2 >= 0.15 else "small" if f2 >= 0.02 else "none"


def effect_label_q2(q2: float) -> str:
    if not np.isfinite(q2):
        return "-"
    return "large" if q2 >= 0.50 else "medium" if q2 >= 0.25 else "small" if q2 > 0 else "no relevance"


def setup_matplotlib() -> None:
    plt.rcParams.update({
        "font.family": "DejaVu Sans",
        "font.size": 9,
        "axes.edgecolor": C_GRID,
        "axes.labelcolor": C_TEXT2,
        "axes.titlecolor": C_TEXT,
        "axes.titlesize": 10,
        "axes.titleweight": "bold",
        "axes.grid": True,
        "grid.color": C_GRID,
        "grid.linewidth": 0.6,
        "xtick.color": C_TEXT2,
        "ytick.color": C_TEXT2,
        "axes.spines.top": False,
        "axes.spines.right": False,
        "figure.facecolor": "white",
        "axes.facecolor": "white",
        "savefig.facecolor": "white",
        "legend.frameon": False,
    })


# =============================================================================
# 3. PLS-SEM ENGINE
# =============================================================================

class ModelSpec:
    """Measurement blocks + structural paths."""

    def __init__(self, blocks: "OrderedDict[str, list[str]]", paths: list[tuple[str, str]],
                 modes: dict | None = None):
        self.lv = list(blocks.keys())
        self.L = len(self.lv)
        self.blocks = OrderedDict((k, list(v)) for k, v in blocks.items())
        self.indicators = [i for v in self.blocks.values() for i in v]
        pos, self.block_idx = 0, []
        for v in self.blocks.values():
            self.block_idx.append(np.arange(pos, pos + len(v)))
            pos += len(v)
        self.paths = list(paths)
        self.lv_index = {n: i for i, n in enumerate(self.lv)}
        self.pred = [[] for _ in range(self.L)]
        self.succ = [[] for _ in range(self.L)]
        for s, t in self.paths:
            self.pred[self.lv_index[t]].append(self.lv_index[s])
            self.succ[self.lv_index[s]].append(self.lv_index[t])
        self.endog = [j for j in range(self.L) if self.pred[j]]
        self.exog = [j for j in range(self.L) if not self.pred[j]]
        self.modes = [(modes or {}).get(n, "A") for n in self.lv]
        self.ind_lv = np.concatenate([[j] * len(b) for j, b in enumerate(self.block_idx)]).astype(int)

    def without_path(self, s: str, t: str) -> "ModelSpec":
        return ModelSpec(self.blocks, [p for p in self.paths if p != (s, t)],
                         dict(zip(self.lv, self.modes)))

    def with_blocks(self, blocks) -> "ModelSpec":
        return ModelSpec(blocks, self.paths, dict(zip(self.lv, self.modes)))

    def topo_order(self) -> list[int]:
        order, done = [], set()
        while len(order) < self.L:
            for j in range(self.L):
                if j not in done and all(p in done for p in self.pred[j]):
                    order.append(j)
                    done.add(j)
        return order


@dataclass
class PLSResult:
    spec: ModelSpec
    W: list                 # outer weights per block (scaled: LV variance 1)
    Y: np.ndarray           # LV scores (standardised)
    C: np.ndarray           # LV correlation matrix
    B: np.ndarray           # path matrix B[i, j] = i -> j
    R2: np.ndarray
    loadings: np.ndarray    # per indicator (own construct)
    weights: np.ndarray     # per indicator
    iterations: int
    converged: bool
    Z: np.ndarray = field(repr=False, default=None)  # standardised data used

    @property
    def total(self) -> np.ndarray:
        L = self.spec.L
        return self.B @ np.linalg.inv(np.eye(L) - self.B)

    @property
    def indirect(self) -> np.ndarray:
        return self.total - self.B

    def adj_r2(self) -> np.ndarray:
        n = self.Y.shape[0]
        out = np.full(self.spec.L, np.nan)
        for j in self.spec.endog:
            k = len(self.spec.pred[j])
            out[j] = 1 - (1 - self.R2[j]) * (n - 1) / (n - k - 1)
        return out


def _inner_weights(C: np.ndarray, spec: ModelSpec, scheme: str) -> np.ndarray:
    L = spec.L
    E = np.zeros((L, L))
    for j in range(L):
        if scheme == "path":
            P = spec.pred[j]
            if P:
                E[P, j] = np.linalg.lstsq(C[np.ix_(P, P)], C[P, j], rcond=None)[0]
            for k in spec.succ[j]:
                E[k, j] = C[k, j]
        else:
            for k in spec.pred[j] + spec.succ[j]:
                E[k, j] = C[k, j] if scheme == "factorial" else np.sign(C[k, j])
    return E


def pls_core(Z: np.ndarray, spec: ModelSpec, scheme: str = INNER_WEIGHTING,
             max_iter: int = MAX_ITER, tol: float = TOLERANCE) -> PLSResult:
    """PLS path modelling (Lohmoeller algorithm) on standardised data Z."""
    n = Z.shape[0]
    W = []
    for b in spec.block_idx:
        w = np.ones(len(b))
        y = Z[:, b] @ w
        sd = y.std()
        W.append(w / (sd if sd > 0 else 1.0))
    converged, it = False, 0
    for it in range(1, max_iter + 1):
        Y = np.column_stack([Z[:, b] @ w for b, w in zip(spec.block_idx, W)])
        C = (Y.T @ Y) / n
        E = _inner_weights(C, spec, scheme)
        Zt = Y @ E
        W_new = []
        for j, b in enumerate(spec.block_idx):
            Xb = Z[:, b]
            if spec.modes[j] == "B" and len(b) > 1:
                w = np.linalg.lstsq(Xb, Zt[:, j], rcond=None)[0]
            else:
                w = Xb.T @ Zt[:, j] / n
            y = Xb @ w
            sd = y.std()
            W_new.append(w / (sd if sd > 0 else 1.0))
        delta = max(np.max(np.abs(a - b)) for a, b in zip(W_new, W))
        W = W_new
        if delta < tol:
            converged = True
            break
    Y = np.column_stack([Z[:, b] @ w for b, w in zip(spec.block_idx, W)])
    Y = Y / np.where(Y.std(axis=0) == 0, 1, Y.std(axis=0))
    C = (Y.T @ Y) / n
    B = np.zeros((spec.L, spec.L))
    R2 = np.full(spec.L, np.nan)
    for j in spec.endog:
        P = spec.pred[j]
        beta = np.linalg.lstsq(C[np.ix_(P, P)], C[P, j], rcond=None)[0]
        B[P, j] = beta
        R2[j] = float(beta @ C[P, j])
    loadings = np.zeros(len(spec.indicators))
    weights = np.zeros(len(spec.indicators))
    for j, b in enumerate(spec.block_idx):
        loadings[b] = Z[:, b].T @ Y[:, j] / n
        weights[b] = W[j]
    return PLSResult(spec, W, Y, C, B, R2, loadings, weights, it, converged, Z)


def fit_pls(X: np.ndarray, spec: ModelSpec, scheme: str = INNER_WEIGHTING) -> PLSResult:
    Z, _, _ = standardize(X)
    return pls_core(Z, spec, scheme)


def cross_loadings(res: PLSResult) -> np.ndarray:
    n = res.Z.shape[0]
    return res.Z.T @ res.Y / n


def htmt_matrix(R: np.ndarray, spec: ModelSpec, kind: str = "htmt") -> np.ndarray:
    """HTMT (arithmetic means) or HTMT2 (geometric means) from item correlations."""
    A = np.abs(R)
    L = spec.L
    out = np.full((L, L), np.nan)

    def mean_(v):
        v = np.clip(v, 1e-12, None)
        return float(np.exp(np.mean(np.log(v)))) if kind == "htmt2" else float(np.mean(v))

    mono = []
    for b in spec.block_idx:
        if len(b) < 2:
            mono.append(np.nan)
            continue
        sub = A[np.ix_(b, b)]
        mono.append(mean_(sub[np.triu_indices(len(b), 1)]))
    for i in range(L):
        for j in range(i):
            het = mean_(A[np.ix_(spec.block_idx[i], spec.block_idx[j])].ravel())
            out[i, j] = het / math.sqrt(mono[i] * mono[j]) if np.isfinite(mono[i] * mono[j]) else np.nan
    return out


def reliability(res: PLSResult, R: np.ndarray) -> pd.DataFrame:
    rows = []
    for j, b in enumerate(res.spec.block_idx):
        lam = res.loadings[b]
        k = len(b)
        Rb = R[np.ix_(b, b)]
        alpha = k / (k - 1) * (1 - k / Rb.sum()) if k > 1 else np.nan
        w = res.W[j]
        if k > 1:
            num = w @ (Rb - np.diag(np.diag(Rb))) @ w
            ww = np.outer(w, w)
            den = w @ (ww - np.diag(np.diag(ww))) @ w
            rho_a = (w @ w) ** 2 * num / den
        else:
            rho_a = 1.0
        rho_c = lam.sum() ** 2 / (lam.sum() ** 2 + (1 - lam ** 2).sum())
        ave = float(np.mean(lam ** 2))
        rows.append([res.spec.lv[j], k, alpha, rho_a, rho_c, ave, math.sqrt(ave)])
    return pd.DataFrame(rows, columns=["Construct", "Items", "Cronbach_alpha", "rho_A",
                                       "Composite_reliability_rho_C", "AVE", "sqrt_AVE"])


def outer_vif(R: np.ndarray, spec: ModelSpec) -> np.ndarray:
    out = np.ones(len(spec.indicators))
    for b in spec.block_idx:
        if len(b) > 1:
            out[b] = np.diag(np.linalg.pinv(R[np.ix_(b, b)]))
    return out


def inner_vif(res: PLSResult) -> dict:
    out = {}
    for j in res.spec.endog:
        P = res.spec.pred[j]
        v = np.diag(np.linalg.pinv(res.C[np.ix_(P, P)])) if len(P) > 1 else np.ones(1)
        for p, val in zip(P, v):
            out[(p, j)] = float(val)
    return out


def r2_from_corr(C: np.ndarray, P: list[int], j: int) -> float:
    if not P:
        return 0.0
    beta = np.linalg.lstsq(C[np.ix_(P, P)], C[P, j], rcond=None)[0]
    return float(beta @ C[P, j])


def f2_values(res: PLSResult) -> dict:
    out = {}
    for j in res.spec.endog:
        P = res.spec.pred[j]
        for p in P:
            r2_ex = r2_from_corr(res.C, [q for q in P if q != p], j)
            out[(p, j)] = (res.R2[j] - r2_ex) / (1 - res.R2[j])
    return out


def plsc_correct(res: PLSResult, R: np.ndarray):
    """Consistent PLS (Dijkstra & Henseler, 2015)."""
    spec = res.spec
    rho = np.ones(spec.L)
    lam = np.zeros(len(spec.indicators))
    for j, b in enumerate(spec.block_idx):
        w = res.W[j]
        if len(b) > 1:
            Rb = R[np.ix_(b, b)]
            num = w @ (Rb - np.diag(np.diag(Rb))) @ w
            ww = np.outer(w, w)
            den = w @ (ww - np.diag(np.diag(ww))) @ w
            c2 = num / den
            rho[j] = (w @ w) ** 2 * c2
            lam[b] = math.sqrt(max(c2, 0)) * w
        else:
            lam[b] = 1.0
    Cc = res.C / np.sqrt(np.outer(rho, rho))
    np.fill_diagonal(Cc, 1.0)
    B = np.zeros((spec.L, spec.L))
    R2 = np.full(spec.L, np.nan)
    for j in spec.endog:
        P = spec.pred[j]
        beta = np.linalg.lstsq(Cc[np.ix_(P, P)], Cc[P, j], rcond=None)[0]
        B[P, j] = beta
        R2[j] = float(beta @ Cc[P, j])
    return {"B": B, "R2": R2, "C": Cc, "rho_A": rho, "loadings": lam}


# ---- model fit ---------------------------------------------------------------

def implied_lv_corr(res: PLSResult, C: np.ndarray | None = None, B: np.ndarray | None = None,
                    R2: np.ndarray | None = None) -> np.ndarray:
    """Model-implied construct correlations of the estimated (structural) model."""
    spec = res.spec
    C = res.C if C is None else C
    B = res.B if B is None else B
    R2 = res.R2 if R2 is None else R2
    L = spec.L
    Psi = np.zeros((L, L))
    ex = spec.exog
    Psi[np.ix_(ex, ex)] = C[np.ix_(ex, ex)]
    for j in spec.endog:
        Psi[j, j] = 1 - R2[j]
    A = np.linalg.inv(np.eye(L) - B.T)
    return A @ Psi @ A.T


def implied_indicator_corr(lam: np.ndarray, ind_lv: np.ndarray, Rlv: np.ndarray) -> np.ndarray:
    S = np.outer(lam, lam) * Rlv[np.ix_(ind_lv, ind_lv)]
    np.fill_diagonal(S, 1.0)
    return S


def discrepancies(S: np.ndarray, Sig: np.ndarray, n: int) -> dict:
    p = S.shape[0]
    D = S - Sig
    tri = np.tril_indices(p)
    srmr = math.sqrt(np.mean(D[tri] ** 2))
    duls = 0.5 * np.sum(D ** 2)
    out = {"SRMR": srmr, "d_ULS": duls, "d_G": np.nan, "Chi_square": np.nan, "NFI": np.nan}
    try:
        ev_s = np.linalg.eigvalsh(S)
        ev_m = np.linalg.eigvalsh(Sig)
        if ev_s.min() > 0 and ev_m.min() > 0:
            ev = np.linalg.eigvals(np.linalg.solve(S, Sig)).real
            ev = ev[ev > 0]
            out["d_G"] = 0.5 * np.sum(np.log(ev) ** 2)
            _, ld_m = np.linalg.slogdet(Sig)
            _, ld_s = np.linalg.slogdet(S)
            F = ld_m + np.trace(S @ np.linalg.inv(Sig)) - ld_s - p
            chi = (n - 1) * F
            chi0 = (n - 1) * (-ld_s)
            out["Chi_square"] = chi
            out["NFI"] = 1 - chi / chi0
    except np.linalg.LinAlgError:
        pass
    return out


def model_fit(res: PLSResult) -> dict:
    Z = res.Z
    n = Z.shape[0]
    S = (Z.T @ Z) / n
    spec = res.spec
    sat = implied_indicator_corr(res.loadings, spec.ind_lv, res.C)
    est = implied_indicator_corr(res.loadings, spec.ind_lv, implied_lv_corr(res))
    out = {"saturated": discrepancies(S, sat, n), "estimated": discrepancies(S, est, n),
           "Sigma_sat": sat, "Sigma_est": est, "S": S}
    # RMS_theta: correlations among outer-model residuals
    E = Z - res.Y[:, spec.ind_lv] * res.loadings
    Re = corr(E)
    off = Re[~np.eye(Re.shape[0], dtype=bool)]
    out["RMS_theta"] = math.sqrt(np.mean(off ** 2))
    comm = np.mean(res.loadings ** 2)
    out["GoF"] = math.sqrt(comm * np.nanmean(res.R2[spec.endog]))
    return out


# ---- blindfolding, PLSpredict --------------------------------------------------

def blindfolding(Z: np.ndarray, spec: ModelSpec, D: int = BLINDFOLD_D,
                 targets: list[int] | None = None, scheme: str = INNER_WEIGHTING) -> dict:
    """Cross-validated redundancy Q2 (omission distance D)."""
    n = Z.shape[0]
    out = {}
    for j in (targets if targets is not None else spec.endog):
        b = spec.block_idx[j]
        k = len(b)
        # column-wise numbering avoids omitting whole indicators when k is a multiple of D
        cell = (np.arange(k)[None, :] * n + np.arange(n)[:, None])
        sse = sso = 0.0
        for d in range(D):
            mask = (cell % D) == d
            Zm = Z.copy()
            sub = Zm[:, b]
            sub[mask] = 0.0
            Zm[:, b] = sub
            Zs, mm, sm = standardize(Zm)
            r = pls_core(Zs, spec, scheme)
            yhat = r.Y[:, spec.pred[j]] @ r.B[spec.pred[j], j]
            xhat = np.outer(yhat, r.loadings[b]) * sm[b] + mm[b]
            orig = Z[:, b]
            sse += np.sum((orig[mask] - xhat[mask]) ** 2)
            sso += np.sum(orig[mask] ** 2)
        out[j] = 1 - sse / sso
    return out


def pls_predict(X: np.ndarray, spec: ModelSpec, folds: int = PREDICT_FOLDS,
                repeats: int = PREDICT_REPEATS, seed: int = SEED, scheme: str = INNER_WEIGHTING):
    n, _ = X.shape
    rng = np.random.default_rng(seed)
    endo_ind = np.concatenate([spec.block_idx[j] for j in spec.endog])
    exo_ind = np.concatenate([spec.block_idx[j] for j in spec.exog])
    k = len(endo_ind)
    se_pls = np.zeros(k); ae_pls = np.zeros(k)
    se_lm = np.zeros(k); ae_lm = np.zeros(k)
    se_naive = np.zeros(k)
    lv_se = np.zeros(spec.L); lv_sn = np.zeros(spec.L)
    count = 0
    for _ in range(repeats):
        perm = rng.permutation(n)
        for f in range(folds):
            test = perm[f::folds]
            train = np.setdiff1d(perm, test)
            Xtr, Xte = X[train], X[test]
            Ztr, m, s = standardize(Xtr)
            r = pls_core(Ztr, spec, scheme)
            Zte = (Xte - m) / s
            Yte = np.column_stack([Zte[:, b] @ w for b, w in zip(spec.block_idx, r.W)])
            T = r.total
            Yhat = np.zeros_like(Yte)
            for j in spec.endog:
                Yhat[:, j] = Yte[:, spec.exog] @ T[spec.exog, j]
            pred = np.zeros((len(test), k))
            for c, idx in enumerate(endo_ind):
                j = spec.ind_lv[idx]
                pred[:, c] = Yhat[:, j] * r.loadings[idx] * s[idx] + m[idx]
            actual = Xte[:, endo_ind]
            se_pls += np.sum((actual - pred) ** 2, axis=0)
            ae_pls += np.sum(np.abs(actual - pred), axis=0)
            se_naive += np.sum((actual - Xtr[:, endo_ind].mean(axis=0)) ** 2, axis=0)
            A = np.column_stack([np.ones(len(train)), Xtr[:, exo_ind]])
            Bm = np.linalg.lstsq(A, Xtr[:, endo_ind], rcond=None)[0]
            plm = np.column_stack([np.ones(len(test)), Xte[:, exo_ind]]) @ Bm
            se_lm += np.sum((actual - plm) ** 2, axis=0)
            ae_lm += np.sum(np.abs(actual - plm), axis=0)
            for j in spec.endog:
                lv_se[j] += np.sum((Yte[:, j] - Yhat[:, j]) ** 2)
                lv_sn[j] += np.sum(Yte[:, j] ** 2)
            count += len(test)
    ind = pd.DataFrame({
        "Indicator": [spec.indicators[i] for i in endo_ind],
        "Construct": [spec.lv[spec.ind_lv[i]] for i in endo_ind],
        "Q2_predict": 1 - se_pls / se_naive,
        "PLS_RMSE": np.sqrt(se_pls / count), "PLS_MAE": ae_pls / count,
        "LM_RMSE": np.sqrt(se_lm / count), "LM_MAE": ae_lm / count,
    })
    ind["RMSE_diff_PLS_minus_LM"] = ind["PLS_RMSE"] - ind["LM_RMSE"]
    ind["MAE_diff_PLS_minus_LM"] = ind["PLS_MAE"] - ind["LM_MAE"]
    lv = pd.DataFrame({
        "Construct": [spec.lv[j] for j in spec.endog],
        "Q2_predict": [1 - lv_se[j] / lv_sn[j] for j in spec.endog],
        "RMSE": [math.sqrt(lv_se[j] / count) for j in spec.endog],
    })
    return ind, lv


# ---- bootstrap ---------------------------------------------------------------

def bootstrap(X: np.ndarray, spec: ModelSpec, n_boot: int, seed: int,
              extractor, scheme: str = INNER_WEIGHTING, label: str = "") -> np.ndarray:
    rng = np.random.default_rng(seed)
    n = X.shape[0]
    out = []
    t0 = time.time()
    step = max(1, n_boot // 4)
    for bidx in range(n_boot):
        idx = rng.integers(0, n, n)
        Z, _, _ = standardize(X[idx])
        try:
            r = pls_core(Z, spec, scheme)
            out.append(extractor(r))
        except np.linalg.LinAlgError:
            continue
        if (bidx + 1) % step == 0:
            log(f"   bootstrap {label} {bidx + 1}/{n_boot} ({time.time() - t0:.0f}s)")
    return np.array(out)


def jackknife(X: np.ndarray, spec: ModelSpec, extractor, scheme: str = INNER_WEIGHTING) -> np.ndarray:
    n = X.shape[0]
    out = []
    for i in range(n):
        Z, _, _ = standardize(np.delete(X, i, axis=0))
        out.append(extractor(pls_core(Z, spec, scheme)))
    return np.array(out)


def ci_table(est: np.ndarray, boot: np.ndarray, jack: np.ndarray | None, n_obs: int,
             level: float = 0.95) -> pd.DataFrame:
    a = (1 - level) / 2
    mean = np.nanmean(boot, axis=0)
    sd = np.nanstd(boot, axis=0, ddof=1)
    t = est / np.where(sd == 0, np.nan, sd)
    p = np.array([p_from_t(v, n_obs - 1) for v in t])
    lo, hi = np.nanpercentile(boot, 100 * a, axis=0), np.nanpercentile(boot, 100 * (1 - a), axis=0)
    # bias-corrected
    prop = np.clip(np.mean(boot < est, axis=0), 1e-6, 1 - 1e-6)
    z0 = stats.norm.ppf(prop)
    za = stats.norm.ppf([a, 1 - a])
    bc_lo = np.array([np.nanpercentile(boot[:, i], 100 * stats.norm.cdf(2 * z0[i] + za[0]))
                      for i in range(boot.shape[1])])
    bc_hi = np.array([np.nanpercentile(boot[:, i], 100 * stats.norm.cdf(2 * z0[i] + za[1]))
                      for i in range(boot.shape[1])])
    out = pd.DataFrame({"Original_sample": est, "Sample_mean": mean, "Std_dev": sd,
                        "t_statistic": t, "p_value": p,
                        "CI_2.5_percentile": lo, "CI_97.5_percentile": hi,
                        "CI_2.5_BC": bc_lo, "CI_97.5_BC": bc_hi})
    if jack is not None:
        jm = jack.mean(axis=0)
        num = np.sum((jm - jack) ** 3, axis=0)
        den = 6 * np.sum((jm - jack) ** 2, axis=0) ** 1.5
        acc = np.where(den == 0, 0, num / den)
        lo_q = stats.norm.cdf(z0 + (z0 + za[0]) / (1 - acc * (z0 + za[0])))
        hi_q = stats.norm.cdf(z0 + (z0 + za[1]) / (1 - acc * (z0 + za[1])))
        out["CI_2.5_BCa"] = [np.nanpercentile(boot[:, i], 100 * lo_q[i]) for i in range(boot.shape[1])]
        out["CI_97.5_BCa"] = [np.nanpercentile(boot[:, i], 100 * hi_q[i]) for i in range(boot.shape[1])]
    return out


# ---- covariance-based SEM with semopy ------------------------------------------

def semopy_description(blocks, paths, structural: bool) -> str:
    desc = "\n".join(f"{c} =~ " + " + ".join(v) for c, v in blocks.items())
    if structural:
        targets = list(OrderedDict.fromkeys(t for _, t in paths))
        desc += "\n" + "\n".join(f"{t} ~ " + " + ".join(s for s, tt in paths if tt == t) for t in targets)
    return desc


def cb_fit_indices(model) -> "OrderedDict[str, float]":
    """Complete set of CB-SEM fit indices from a fitted semopy model.

    Chi-square, df and the baseline (independence) model come from semopy.
    GFI, AGFI, PGFI, SRMR, RMR, IFI, RFI, PNFI, PCFI, RMSEA CI and PCLOSE are
    computed from the sample (S) and model-implied (Sigma) covariance matrices
    with the standard formulas (Joreskog & Sorbom; Bentler; Bollen; Mulaik et al.).
    Note: semopy's own "GFI" equals 1 - chi2/chi2_baseline (i.e. the NFI), so the
    Joreskog-Sorbom GFI is recomputed here.
    """
    import semopy
    S = np.asarray(model.mx_cov, dtype=float)
    Sig = np.asarray(model.calc_sigma()[0], dtype=float)
    n = int(model.n_samples)
    p = S.shape[0]
    st = semopy.calc_stats(model).T.iloc[:, 0]
    chi2, df = float(st["chi2"]), float(st["DoF"])
    chi0, df0 = float(st["chi2 Baseline"]), float(st["DoF Baseline"])
    n_par = len(model.param_vals)
    Si = np.linalg.inv(Sig)
    SiS = Si @ S
    I = np.eye(p)
    gfi = 1 - np.trace((SiS - I) @ (SiS - I)) / np.trace(SiS @ SiS)
    nmom = p * (p + 1) / 2
    agfi = 1 - nmom / df * (1 - gfi) if df > 0 else np.nan
    pgfi = df / nmom * gfi
    d = np.sqrt(np.diag(S))
    tri = np.tril_indices(p)
    srmr = math.sqrt(np.mean(((S - Sig) / np.outer(d, d))[tri] ** 2))
    rmr = math.sqrt(np.mean((S - Sig)[tri] ** 2))
    nfi = 1 - chi2 / chi0
    rfi = 1 - (chi2 / df) / (chi0 / df0) if df > 0 else np.nan
    ifi = (chi0 - chi2) / (chi0 - df)
    tli = (chi0 / df0 - chi2 / df) / (chi0 / df0 - 1) if df > 0 else np.nan
    cfi = 1 - max(chi2 - df, 0) / max(chi0 - df0, chi2 - df, 1e-12)
    rmsea = math.sqrt(max(chi2 - df, 0) / (df * (n - 1))) if df > 0 else np.nan

    def nc_bound(q):
        """Noncentrality lambda with P(chi2_df(lambda) <= observed) = q."""
        from scipy.optimize import brentq
        f = lambda lam: stats.ncx2.cdf(chi2, df, lam) - q  # noqa: E731
        if stats.chi2.cdf(chi2, df) <= q:
            return 0.0
        hi = max(chi2, 10.0)
        while f(hi) > 0:
            hi *= 2
        return brentq(f, 1e-10, hi)

    rmsea_lo = math.sqrt(nc_bound(0.95) / (df * (n - 1)))
    rmsea_hi = math.sqrt(nc_bound(0.05) / (df * (n - 1)))
    pclose = float(stats.ncx2.sf(chi2, df, 0.05 ** 2 * df * (n - 1)))
    return OrderedDict([
        ("Chi-square", chi2), ("df", df), ("p-value", float(stats.chi2.sf(chi2, df))),
        ("X2/df", chi2 / df), ("GFI", gfi), ("AGFI", agfi), ("PGFI", pgfi), ("RMR", rmr),
        ("SRMR", srmr), ("NFI", nfi), ("RFI", rfi), ("IFI", ifi), ("TLI", tli), ("CFI", cfi),
        ("PNFI", df / df0 * nfi), ("PCFI", df / df0 * cfi), ("RMSEA", rmsea),
        ("RMSEA 90% CI low", rmsea_lo), ("RMSEA 90% CI high", rmsea_hi), ("PCLOSE", pclose),
        ("AIC (semopy)", float(st["AIC"])), ("BIC (semopy)", float(st["BIC"])),
        ("Baseline Chi-square", chi0), ("Baseline df", df0), ("Free parameters", n_par),
        ("N", n)])


def semopy_analysis(items: pd.DataFrame, blocks, paths, folder: Path):
    """CFA + structural CB-SEM with semopy, estimator robustness and reports."""
    import semopy
    out = OrderedDict()
    models = OrderedDict()
    for label, structural in (("Measurement model (CFA)", False), ("Structural model", True)):
        m = semopy.Model(semopy_description(blocks, paths, structural))
        r = m.fit(items, obj="MLW")
        models[label] = m
        out[label] = {"fit": cb_fit_indices(m), "estimates": m.inspect(std_est=True),
                      "converged": "successful" in str(r).lower()}
    # estimator robustness (structural paths under alternative discrepancy functions)
    est_rows = OrderedDict()
    for obj in ("MLW", "ULS", "GLS"):
        try:
            m = semopy.Model(semopy_description(blocks, paths, True))
            m.fit(items, obj=obj)
            ins = m.inspect(std_est=True)
            ins = ins[(ins["op"] == "~") & ins["lval"].isin(list(blocks)) & ins["rval"].isin(list(blocks))]
            est_rows[obj] = {f"{rv} -> {lv}": est for lv, rv, est in
                             zip(ins["lval"], ins["rval"], ins["Est. Std"])}
        except Exception as exc:  # noqa: BLE001
            log(f"   semopy {obj} failed: {exc}")
    out["estimators"] = pd.DataFrame(est_rows)
    # semopy HTML report and path diagram (need graphviz; skipped silently if absent)
    import logging
    logging.getLogger().setLevel(logging.ERROR)  # silence semopy's graphviz warning
    try:
        semopy.semplot(models["Structural model"], str(folder / "semopy_structural_model.png"),
                       plot_covs=True, std_ests=True)
        sp_png = folder / "semopy_structural_model.png"
        out["semplot"] = sp_png if sp_png.exists() else None
    except Exception:  # noqa: BLE001
        out["semplot"] = None
    try:
        semopy.report(models["Structural model"], str(folder / "semopy_report"))
        out["report"] = folder / "semopy_report" if (folder / "semopy_report").exists() else None
    except Exception:  # noqa: BLE001
        out["report"] = None
    return out


def fit_indices_table(fi: dict, label: str) -> pd.DataFrame:
    """Table in the layout Category / The fit indices / ATV / Result / Decision."""
    spec_ = [("Absolute", "X2/df", "1-5", lambda v: 1 <= v <= 5),
             ("Absolute", "GFI", "> 0.9", lambda v: v > 0.9),
             ("Absolute", "SRMR", "< 0.08", lambda v: v < 0.08),
             ("Relative", "IFI", "> 0.9", lambda v: v > 0.9),
             ("Relative", "NFI", "> 0.9", lambda v: v > 0.9),
             ("Parsimonious", "PGFI", "> 0.50", lambda v: v > 0.5),
             ("Parsimonious", "PNFI", "> 0.50", lambda v: v > 0.5),
             ("Noncentrality", "RMSEA", "< 0.1", lambda v: v < 0.1),
             ("Noncentrality", "CFI", "> 0.9", lambda v: v > 0.9)]
    cols = ["Category"] + [c for c, _, _, _ in spec_]
    rows = [["The fit indices"] + [n for _, n, _, _ in spec_],
            ["ATV*"] + [a for _, _, a, _ in spec_],
            [f"Result ({label})"] + [f"{fi[n]:.3f}" for _, n, _, _ in spec_],
            ["Decision"] + ["Acceptable" if ok(fi[n]) else "Not acceptable" for _, n, _, ok in spec_]]
    return pd.DataFrame(rows, columns=cols)


def boxed_table(df: pd.DataFrame, header: bool = False) -> str:
    """Box-drawing text table (tabulate 'fancy_grid' look) without extra dependencies."""
    data = ([list(map(str, df.columns))] if header else []) + [[str(v) for v in r] for r in df.values]
    w = [max(len(r[i]) for r in data) for i in range(len(data[0]))]

    def line(l, m, r):
        return l + m.join("─" * (x + 2) for x in w) + r

    out = ["╒" + "╤".join("═" * (x + 2) for x in w) + "╕"]
    for k, r in enumerate(data):
        out.append("│" + "│".join(f" {c:<{x}} " for c, x in zip(r, w)) + "│")
        if k == 0:
            out.append("╞" + "╪".join("═" * (x + 2) for x in w) + "╡")
        elif k < len(data) - 1:
            out.append(line("├", "┼", "┤"))
    out.append("╘" + "╧".join("═" * (x + 2) for x in w) + "╛")
    return "\n".join(out)


# =============================================================================
# 4. DATA LOADING AND SCREENING
# =============================================================================

def locate_input() -> Path:
    if len(sys.argv) > 1:
        p = Path(sys.argv[1]).expanduser()
        if p.exists():
            return p
        raise FileNotFoundError(f"File not found: {p}")
    home = Path.home()
    candidates = [home / "Desktop" / INPUT_FILE_NAME,
                  home / "OneDrive" / "Desktop" / INPUT_FILE_NAME,
                  Path(__file__).resolve().parent / INPUT_FILE_NAME,
                  Path.cwd() / INPUT_FILE_NAME]
    candidates += list(home.glob(f"OneDrive*/Desktop/{INPUT_FILE_NAME}"))
    candidates += list(home.glob(f"OneDrive*/*/{INPUT_FILE_NAME}"))
    for c in candidates:
        if c.exists():
            return c
    raise FileNotFoundError(
        f"Could not find {INPUT_FILE_NAME} on the Desktop. Pass the full path:\n"
        f"    python {Path(__file__).name} \"C:/path/to/{INPUT_FILE_NAME}\"")


def load_data(path: Path):
    raw = pd.read_excel(path, sheet_name=SHEET_NAME)
    raw.columns = [str(c).strip() for c in raw.columns]
    upper = {c.upper(): c for c in raw.columns}
    all_items = [i for c in CONSTRUCTS.values() for i in c["items"]]
    rename = {}
    for name in list(DEMOGRAPHICS) + all_items:
        if name not in raw.columns and name.upper() in upper:
            rename[upper[name.upper()]] = name
    raw = raw.rename(columns=rename)
    missing_cols = [c for c in all_items if c not in raw.columns]
    if missing_cols:
        raise ValueError(f"Columns missing from the Excel file: {missing_cols}")
    raw = raw.dropna(how="all").reset_index(drop=True)
    screening = OrderedDict()
    items = raw[all_items].apply(pd.to_numeric, errors="coerce")
    out_of_range = ((items < SCALE_MIN) | (items > SCALE_MAX)).sum()
    items = items.where((items >= SCALE_MIN) & (items <= SCALE_MAX))
    miss = items.isna().sum()
    screening["missing"] = pd.DataFrame({"Item": all_items, "Missing": miss.values,
                                         "Missing_%": (miss.values / len(items) * 100),
                                         "Out_of_range_set_missing": out_of_range.values})
    n_raw = len(items)
    if MISSING_STRATEGY == "listwise":
        keep = items.notna().all(axis=1)
        raw, items = raw[keep].reset_index(drop=True), items[keep].reset_index(drop=True)
    else:
        items = items.fillna(items.mean())
    sd_row = items.std(axis=1)
    straight = sd_row == 0
    screening["straight_lining"] = pd.DataFrame({
        "Case": np.arange(1, len(items) + 1), "Row_SD_of_items": sd_row.values,
        "Straight_liner": straight.values})
    log(f"Loaded {n_raw} cases from {path.name}; analysed n = {len(items)}; "
        f"missing cells = {int(miss.sum())}; straight-liners = {int(straight.sum())}")
    demo = raw[[c for c in DEMOGRAPHICS if c in raw.columns]].copy()
    return raw, items.astype(float), demo, screening


def mahalanobis(items: pd.DataFrame) -> pd.DataFrame:
    X = items.values
    d = X - X.mean(axis=0)
    S_inv = np.linalg.pinv(np.cov(X, rowvar=False))
    d2 = np.einsum("ij,jk,ik->i", d, S_inv, d)
    p = stats.chi2.sf(d2, X.shape[1])
    return pd.DataFrame({"Case": np.arange(1, len(X) + 1), "Mahalanobis_D2": d2,
                         "p_value": p, "Outlier_p<0.001": p < 0.001})


# =============================================================================
# 5. DESCRIPTIVE, NORMALITY AND DIMENSIONALITY
# =============================================================================

def demographic_table(demo: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for var, labels in DEMOGRAPHICS.items():
        if var not in demo.columns:
            continue
        col = demo[var]
        num = pd.to_numeric(col, errors="coerce")
        series = num.map(labels).fillna(col.astype(str)) if num.notna().any() else col.astype(str)
        counts = series.value_counts(dropna=True)
        order = [v for v in labels.values() if v in counts.index] + \
                [v for v in counts.index if v not in labels.values()]
        for cat in order:
            rows.append([var, cat, int(counts[cat]), counts[cat] / counts.sum() * 100])
        rows.append([var, "Total", int(counts.sum()), 100.0])
    return pd.DataFrame(rows, columns=["Variable", "Category", "Frequency", "Percent"])


def descriptive_items(items: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for c, info in CONSTRUCTS.items():
        for it in info["items"]:
            x = items[it].values
            sw = stats.shapiro(x) if len(x) >= 3 else (np.nan, np.nan)
            z = (x - x.mean()) / x.std(ddof=1)
            ks = stats.kstest(z, "norm")
            rows.append([c, it, len(x), x.mean(), x.std(ddof=1), x.min(), x.max(), np.median(x),
                         stats.skew(x, bias=False), stats.kurtosis(x, bias=False),
                         sw[0], sw[1], ks.statistic, ks.pvalue])
    df = pd.DataFrame(rows, columns=["Construct", "Item", "N", "Mean", "SD", "Min", "Max", "Median",
                                     "Skewness", "Excess_kurtosis", "Shapiro_W", "Shapiro_p",
                                     "KS_D", "KS_p"])
    df["Within_+-1_skew_kurt"] = (df["Skewness"].abs() <= 1) & (df["Excess_kurtosis"].abs() <= 1)
    df["Within_+-2_skew_kurt"] = (df["Skewness"].abs() <= 2) & (df["Excess_kurtosis"].abs() <= 2)
    return df


def construct_means(items: pd.DataFrame, blocks) -> pd.DataFrame:
    return pd.DataFrame({c: items[v].mean(axis=1) for c, v in blocks.items()})


def descriptive_constructs(means: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for c in means.columns:
        x = means[c].values
        sw = stats.shapiro(x)
        ks = stats.kstest((x - x.mean()) / x.std(ddof=1), "norm")
        rows.append([c, CONSTRUCTS[c]["name"] if c in CONSTRUCTS else c, x.mean(), x.std(ddof=1),
                     x.min(), x.max(), stats.skew(x, bias=False), stats.kurtosis(x, bias=False),
                     sw[0], sw[1], ks.statistic, ks.pvalue])
    return pd.DataFrame(rows, columns=["Construct", "Name", "Mean", "SD", "Min", "Max", "Skewness",
                                       "Excess_kurtosis", "Shapiro_W", "Shapiro_p", "KS_D", "KS_p"])


def mardia(X: np.ndarray) -> pd.DataFrame:
    n, p = X.shape
    d = X - X.mean(axis=0)
    S = np.cov(X, rowvar=False, ddof=0)
    G = d @ np.linalg.pinv(S) @ d.T
    b1 = np.sum(G ** 3) / n ** 2
    b2 = np.mean(np.diag(G) ** 2)
    skew_stat = n * b1 / 6
    df = p * (p + 1) * (p + 2) / 6
    kurt_z = (b2 - p * (p + 2)) / math.sqrt(8 * p * (p + 2) / n)
    return pd.DataFrame([
        ["Mardia multivariate skewness", b1, skew_stat, df, stats.chi2.sf(skew_stat, df)],
        ["Mardia multivariate kurtosis", b2, kurt_z, np.nan, 2 * stats.norm.sf(abs(kurt_z))],
    ], columns=["Test", "Coefficient", "Statistic", "df", "p_value"])


def kmo_bartlett(X: np.ndarray, items: list[str]):
    R = corr(X)
    n, p = X.shape
    Rinv = np.linalg.pinv(R)
    Pc = -Rinv / np.sqrt(np.outer(np.diag(Rinv), np.diag(Rinv)))
    np.fill_diagonal(Pc, 0)
    R0 = R.copy()
    np.fill_diagonal(R0, 0)
    kmo = np.sum(R0 ** 2) / (np.sum(R0 ** 2) + np.sum(Pc ** 2))
    msa = np.sum(R0 ** 2, axis=0) / (np.sum(R0 ** 2, axis=0) + np.sum(Pc ** 2, axis=0))
    _, logdet = np.linalg.slogdet(R)
    chi = -(n - 1 - (2 * p + 5) / 6) * logdet
    df = p * (p - 1) / 2
    summary = pd.DataFrame([["KMO (overall)", kmo, "", ""],
                            ["Bartlett chi-square", chi, df, stats.chi2.sf(chi, df)]],
                           columns=["Statistic", "Value", "df", "p_value"])
    return summary, pd.DataFrame({"Item": items, "MSA": msa})


def parallel_analysis(X: np.ndarray, reps: int, seed: int):
    n, p = X.shape
    ev = np.sort(np.linalg.eigvalsh(corr(X)))[::-1]
    rng = np.random.default_rng(seed)
    sims = np.array([np.sort(np.linalg.eigvalsh(corr(rng.standard_normal((n, p)))))[::-1]
                     for _ in range(reps)])
    q95 = np.percentile(sims, 95, axis=0)
    df = pd.DataFrame({"Component": np.arange(1, p + 1), "Observed_eigenvalue": ev,
                       "Random_mean": sims.mean(axis=0), "Random_95th": q95,
                       "Variance_%": ev / p * 100, "Cumulative_%": np.cumsum(ev) / p * 100})
    k = 0
    for o, r in zip(ev, q95):
        if o > r:
            k += 1
        else:
            break
    return df, k


def varimax(L: np.ndarray, gamma: float = 1.0, it: int = 500, tol: float = 1e-8) -> np.ndarray:
    p, k = L.shape
    R = np.eye(k)
    d = 0
    for _ in range(it):
        Lr = L @ R
        u, s, vh = np.linalg.svd(L.T @ (Lr ** 3 - (gamma / p) * Lr @ np.diag(np.sum(Lr ** 2, axis=0))))
        R = u @ vh
        d_old, d = d, np.sum(s)
        if d_old != 0 and d / d_old < 1 + tol:
            break
    return L @ R


def promax(L: np.ndarray, power: int = 4):
    Lv = varimax(L)
    P = np.sign(Lv) * np.abs(Lv) ** power
    U = np.linalg.lstsq(Lv, P, rcond=None)[0]
    U = U / np.sqrt(np.diag(np.linalg.inv(U.T @ U)))
    Lp = Lv @ U
    phi = np.linalg.inv(U.T @ U)
    return Lp, phi


def paf(R: np.ndarray, k: int, it: int = 500, tol: float = 1e-6) -> np.ndarray:
    Rr = R.copy()
    h2 = 1 - 1 / np.diag(np.linalg.pinv(R))
    for _ in range(it):
        np.fill_diagonal(Rr, h2)
        vals, vecs = np.linalg.eigh(Rr)
        idx = np.argsort(vals)[::-1][:k]
        L = vecs[:, idx] * np.sqrt(np.clip(vals[idx], 1e-8, None))
        new = np.clip(np.sum(L ** 2, axis=1), 0.001, 0.999)
        if np.max(np.abs(new - h2)) < tol:
            break
        h2 = new
    return L


def efa_table(X: np.ndarray, items: list[str], k: int, rotation: str) -> pd.DataFrame:
    L = paf(corr(X), k)
    Lr = varimax(L) if rotation == "varimax" else promax(L)[0]
    Lr = Lr * np.sign(Lr.sum(axis=0))
    df = pd.DataFrame(Lr, columns=[f"F{i + 1}" for i in range(k)])
    df.insert(0, "Item", items)
    df.insert(1, "Theoretical_construct", [next(c for c, v in CONSTRUCTS.items() if it in v["items"])
                                           for it in items])
    df["Primary_factor"] = [f"F{i + 1}" for i in np.argmax(np.abs(Lr), axis=1)]
    df["Primary_loading"] = np.max(np.abs(Lr), axis=1)
    srt = np.sort(np.abs(Lr), axis=1)
    df["Cross_loading_gap"] = srt[:, -1] - (srt[:, -2] if k > 1 else 0)
    df["Communality"] = np.sum(L ** 2, axis=1)
    return df


# =============================================================================
# 6. ROBUSTNESS TESTS
# =============================================================================

def copula_test(res: PLSResult, n_boot: int, seed: int) -> pd.DataFrame:
    spec = res.spec
    Y = res.Y
    n = Y.shape[0]
    rng = np.random.default_rng(seed)
    rows = []

    def copula(x):
        r = stats.rankdata(x) / (n + 1)
        return stats.norm.ppf(r)

    for j in spec.endog:
        P = spec.pred[j]
        cops = {p: copula(Y[:, p]) for p in P}
        combos = [[p] for p in P] + ([P] if len(P) > 1 else [])
        for combo in combos:
            Xm = np.column_stack([Y[:, P]] + [cops[p] for p in combo])
            b, _, _ = ols(Y[:, j], Xm)
            bs = []
            for _ in range(n_boot):
                idx = rng.integers(0, n, n)
                bs.append(ols(Y[idx, j], Xm[idx])[0])
            bs = np.array(bs)
            for m, p in enumerate(combo):
                pos = 1 + len(P) + m
                se = bs[:, pos].std(ddof=1)
                t = b[pos] / se
                rows.append([spec.lv[j], " + ".join("c(" + spec.lv[q] + ")" for q in combo),
                             f"c({spec.lv[p]})", b[pos], se, t, p_from_t(t, n - 1)])
    df = pd.DataFrame(rows, columns=["Endogenous", "Copula_terms_in_model", "Copula_term",
                                     "Coefficient", "Boot_SE", "t", "p_value"])
    norm = []
    for p in spec.exog + [q for q in spec.endog if spec.succ[q]]:
        x = Y[:, p]
        norm.append([spec.lv[p], stats.shapiro(x)[1], stats.kstest(x, "norm").pvalue])
    nd = pd.DataFrame(norm, columns=["Predictor", "Shapiro_p", "KS_p"])
    nd["Non_normal_(copula_identified)"] = (nd["Shapiro_p"] < 0.05) | (nd["KS_p"] < 0.05)
    return df, nd


def nonlinear_test(res: PLSResult, n_boot: int, seed: int) -> pd.DataFrame:
    spec = res.spec
    Y = res.Y
    n = Y.shape[0]
    rng = np.random.default_rng(seed)
    rows = []
    for j in spec.endog:
        P = spec.pred[j]
        for p in P:
            q = Y[:, p] ** 2
            q = (q - q.mean()) / q.std()
            Xm = np.column_stack([Y[:, P], q])
            b, _, r2_q = ols(Y[:, j], Xm)
            r2_lin = ols(Y[:, j], Y[:, P])[2]
            bs = np.array([ols(Y[idx, j], Xm[idx])[0] for idx in
                           (rng.integers(0, n, n) for _ in range(n_boot))])
            se = bs[:, -1].std(ddof=1)
            t = b[-1] / se
            f2 = (r2_q - r2_lin) / (1 - r2_q)
            rows.append([f"QE({spec.lv[p]}) -> {spec.lv[j]}", b[-1], se, t, p_from_t(t, n - 1), f2])
    return pd.DataFrame(rows, columns=["Quadratic_effect", "Coefficient", "Boot_SE", "t", "p_value", "f2"])


def mga(X: np.ndarray, spec: ModelSpec, groups: np.ndarray, labels: tuple[str, str],
        n_perm: int, n_boot: int, seed: int):
    """Permutation MGA + MICOM steps 2 and 3."""
    rng = np.random.default_rng(seed)
    g1, g2 = groups == 0, groups == 1
    path_ids = [(spec.lv_index[s], spec.lv_index[t]) for s, t in spec.paths]

    def group_fit(mask):
        return fit_pls(X[mask], spec)

    def stats_of(m1, m2):
        r1, r2 = group_fit(m1), group_fit(m2)
        d = np.array([r1.B[i, j] - r2.B[i, j] for i, j in path_ids])
        Zp, _, _ = standardize(X)
        cvals = []
        for b, w1, w2 in zip(spec.block_idx, r1.W, r2.W):
            y1, y2 = Zp[:, b] @ w1, Zp[:, b] @ w2
            cvals.append(np.corrcoef(y1, y2)[0, 1] if len(b) > 1 else 1.0)
        return r1, r2, d, np.array(cvals)

    r1, r2, d_obs, c_obs = stats_of(g1, g2)
    pooled = fit_pls(X, spec)
    Yp = pooled.Y
    mean_d = Yp[g1].mean(axis=0) - Yp[g2].mean(axis=0)
    var_d = np.log(Yp[g1].var(axis=0, ddof=1)) - np.log(Yp[g2].var(axis=0, ddof=1))
    perm_d, perm_c, perm_m, perm_v = [], [], [], []
    for _ in range(n_perm):
        sh = rng.permutation(groups)
        m1, m2 = sh == 0, sh == 1
        try:
            _, _, d, c = stats_of(m1, m2)
        except np.linalg.LinAlgError:
            continue
        perm_d.append(d)
        perm_c.append(c)
        perm_m.append(Yp[m1].mean(axis=0) - Yp[m2].mean(axis=0))
        perm_v.append(np.log(Yp[m1].var(axis=0, ddof=1)) - np.log(Yp[m2].var(axis=0, ddof=1)))
    perm_d, perm_c = np.array(perm_d), np.array(perm_c)
    perm_m, perm_v = np.array(perm_m), np.array(perm_v)
    p_perm = (np.sum(np.abs(perm_d) >= np.abs(d_obs), axis=0) + 1) / (len(perm_d) + 1)

    def boot_p(mask, extractor):
        Xg = X[mask]
        bs = bootstrap(Xg, spec, n_boot, seed, extractor, label="MGA group")
        est = extractor(fit_pls(Xg, spec))
        sd = bs.std(axis=0, ddof=1)
        return [p_from_t(e / s, mask.sum() - 1) for e, s in zip(est, sd)]

    extractor = lambda r: np.array([r.B[i, j] for i, j in path_ids])  # noqa: E731
    p1, p2 = boot_p(g1, extractor), boot_p(g2, extractor)
    paths = pd.DataFrame({
        "Path": [f"{s} -> {t}" for s, t in spec.paths],
        f"beta_{labels[0]}": [r1.B[i, j] for i, j in path_ids], f"p_{labels[0]}": p1,
        f"beta_{labels[1]}": [r2.B[i, j] for i, j in path_ids], f"p_{labels[1]}": p2,
        "Difference": d_obs,
        "Perm_CI_2.5": np.percentile(perm_d, 2.5, axis=0),
        "Perm_CI_97.5": np.percentile(perm_d, 97.5, axis=0),
        "Permutation_p": p_perm})
    micom = pd.DataFrame({
        "Construct": spec.lv,
        "Step2_c": c_obs, "Step2_c_5%_quantile": np.percentile(perm_c, 5, axis=0),
        "Step2_compositional_invariance": c_obs >= np.percentile(perm_c, 5, axis=0),
        "Step3a_mean_diff": mean_d,
        "Step3a_CI_2.5": np.percentile(perm_m, 2.5, axis=0),
        "Step3a_CI_97.5": np.percentile(perm_m, 97.5, axis=0),
        "Step3b_logvar_diff": var_d,
        "Step3b_CI_2.5": np.percentile(perm_v, 2.5, axis=0),
        "Step3b_CI_97.5": np.percentile(perm_v, 97.5, axis=0)})
    micom["Step3a_equal_means"] = (micom["Step3a_mean_diff"] >= micom["Step3a_CI_2.5"]) & \
                                  (micom["Step3a_mean_diff"] <= micom["Step3a_CI_97.5"])
    micom["Step3b_equal_variances"] = (micom["Step3b_logvar_diff"] >= micom["Step3b_CI_2.5"]) & \
                                      (micom["Step3b_logvar_diff"] <= micom["Step3b_CI_97.5"])
    return paths, micom


# =============================================================================
# 7. ITEM PURIFICATION (SENSITIVITY)
# =============================================================================

def purify(X_full: pd.DataFrame, spec: ModelSpec):
    blocks = OrderedDict((k, list(v)) for k, v in spec.blocks.items())
    original = {k: len(v) for k, v in blocks.items()}
    log_rows = []

    def build(bl):
        sp = spec.with_blocks(bl)
        Xm = X_full[sp.indicators].values
        return sp, Xm, fit_pls(Xm, sp), corr(Xm)

    def can_drop(c, bl):
        k = len(bl[c])
        return k > PURIFY_MIN_ITEMS and (original[c] - (k - 1)) / original[c] <= PURIFY_MAX_SHARE

    # Step A - weak loadings
    while True:
        sp, Xm, r, R = build(blocks)
        rel = reliability(r, R).set_index("Construct")
        cands = []
        for c, b in zip(sp.lv, sp.block_idx):
            for i in b:
                lam = r.loadings[i]
                weak = lam < 0.40 or (lam < 0.708 and (rel.loc[c, "AVE"] < 0.50 or
                                                       rel.loc[c, "Composite_reliability_rho_C"] < 0.70))
                if weak and can_drop(c, blocks):
                    cands.append((lam, c, sp.indicators[i]))
        if not cands:
            break
        lam, c, it = min(cands)
        blocks[c].remove(it)
        log_rows.append(["A: low loading", it, c, lam, np.nan])

    # Step B - HTMT driven
    while True:
        sp, Xm, r, R = build(blocks)
        H = htmt_matrix(R, sp)
        hmax = np.nanmax(H)
        if hmax <= PURIFY_HTMT_TARGET:
            break
        i, j = np.unravel_index(np.nanargmax(H), H.shape)
        best = None
        for c in (sp.lv[i], sp.lv[j]):
            if not can_drop(c, blocks):
                continue
            for it in blocks[c]:
                trial = OrderedDict((k, [x for x in v if x != it]) for k, v in blocks.items())
                sp2 = spec.with_blocks(trial)
                h2 = np.nanmax(htmt_matrix(corr(X_full[sp2.indicators].values), sp2))
                if best is None or h2 < best[0]:
                    best = (h2, c, it)
        if best is None or best[0] >= hmax - 1e-4:
            break
        blocks[best[1]].remove(best[2])
        log_rows.append(["B: HTMT", best[2], best[1], np.nan, best[0]])
    log_df = pd.DataFrame(log_rows, columns=["Step", "Item_removed", "Construct",
                                             "Loading_at_removal", "Max_HTMT_after_removal"])
    return blocks, log_df


# =============================================================================
# 8. FIGURES
# =============================================================================

def save(fig, folder: Path, name: str, figs: list) -> None:
    path = folder / f"{name}.png"
    fig.savefig(path, dpi=DPI, bbox_inches="tight")
    plt.close(fig)
    figs.append(path)


def fig_demographics(demo_tab: pd.DataFrame, folder, figs):
    vars_ = [v for v in DEMOGRAPHICS if v in demo_tab["Variable"].unique()]
    if not vars_:
        return
    fig, axes = plt.subplots(1, len(vars_), figsize=(3.4 * len(vars_), 3.2))
    axes = np.atleast_1d(axes)
    for ax, v in zip(axes, vars_):
        d = demo_tab[(demo_tab["Variable"] == v) & (demo_tab["Category"] != "Total")]
        ax.barh(d["Category"], d["Percent"], color=C_SERIES[0], height=0.6)
        for y, (pct, f) in enumerate(zip(d["Percent"], d["Frequency"])):
            ax.text(pct + 1, y, f"{f} ({pct:.1f}%)", va="center", fontsize=8, color=C_TEXT)
        ax.set_title(v)
        ax.set_xlim(0, max(d["Percent"].max() * 1.45, 10))
        ax.set_xlabel("Percent of respondents")
        ax.invert_yaxis()
        ax.grid(axis="y", visible=False)
    fig.suptitle("Demographic profile of the respondents", fontweight="bold", color=C_TEXT)
    fig.tight_layout()
    save(fig, folder, "Fig01_demographics", figs)


def fig_item_means(desc: pd.DataFrame, folder, figs):
    fig, ax = plt.subplots(figsize=(7, 9))
    cons = list(CONSTRUCTS)
    y = np.arange(len(desc))[::-1]
    for k, c in enumerate(cons):
        m = desc["Construct"] == c
        ax.errorbar(desc.loc[m, "Mean"], y[m.values], xerr=desc.loc[m, "SD"], fmt="o",
                    color=C_SERIES[k], ecolor=C_SERIES[k], elinewidth=1.5, capsize=0,
                    markersize=6, label=f"{c} - {CONSTRUCTS[c]['name']}")
    ax.set_yticks(y)
    ax.set_yticklabels(desc["Item"])
    ax.set_xlim(SCALE_MIN - 0.2, SCALE_MAX + 0.2)
    ax.set_xlabel("Mean (+/- 1 SD)")
    ax.set_title("Item means and standard deviations")
    ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.06), ncol=2, fontsize=8)
    fig.tight_layout()
    save(fig, folder, "Fig02_item_means", figs)


def fig_normality(means: pd.DataFrame, folder, figs):
    cols = list(means.columns)
    fig, axes = plt.subplots(2, len(cols), figsize=(2.6 * len(cols), 5))
    for k, c in enumerate(cols):
        x = means[c].values
        ax = axes[0, k]
        ax.hist(x, bins=12, color=C_SERIES[0], edgecolor="white", density=True)
        xs = np.linspace(x.min(), x.max(), 200)
        ax.plot(xs, stats.norm.pdf(xs, x.mean(), x.std(ddof=1)), color=C_TEXT2, lw=2)
        ax.set_title(c)
        ax.grid(axis="x", visible=False)
        ax = axes[1, k]
        (osm, osr), (sl, ic, _) = stats.probplot(x, dist="norm")
        ax.scatter(osm, osr, s=10, color=C_SERIES[0])
        ax.plot(osm, sl * osm + ic, color=C_TEXT2, lw=2)
        ax.set_xlabel("Theoretical quantiles")
        if k == 0:
            ax.set_ylabel("Sample quantiles")
            axes[0, 0].set_ylabel("Density")
    fig.suptitle("Construct score distributions and normal Q-Q plots", fontweight="bold", color=C_TEXT)
    fig.tight_layout()
    save(fig, folder, "Fig03_normality", figs)


def _div_cmap():
    from matplotlib.colors import LinearSegmentedColormap
    return LinearSegmentedColormap.from_list("div", [C_DIV_NEG, C_DIV_MID, C_DIV_POS])


def _seq_cmap():
    from matplotlib.colors import LinearSegmentedColormap
    return LinearSegmentedColormap.from_list("seq", C_BLUE_SEQ)


def fig_item_corr(items: pd.DataFrame, folder, figs):
    R = items.corr().values
    fig, ax = plt.subplots(figsize=(10, 9))
    im = ax.imshow(R, cmap=_div_cmap(), vmin=-1, vmax=1)
    ax.set_xticks(range(len(items.columns)))
    ax.set_yticks(range(len(items.columns)))
    ax.set_xticklabels(items.columns, rotation=90, fontsize=7)
    ax.set_yticklabels(items.columns, fontsize=7)
    ax.grid(False)
    pos = 0
    for info in CONSTRUCTS.values():
        k = len([i for i in info["items"] if i in items.columns])
        ax.add_patch(plt.Rectangle((pos - 0.5, pos - 0.5), k, k, fill=False, ec=C_TEXT, lw=1.2))
        pos += k
    fig.colorbar(im, ax=ax, shrink=0.7, label="Pearson r")
    ax.set_title("Item correlation matrix (boxes = theoretical constructs)")
    save(fig, folder, "Fig04_item_correlations", figs)


def fig_scree(pa: pd.DataFrame, folder, figs):
    fig, ax = plt.subplots(figsize=(7, 4))
    k = min(15, len(pa))
    ax.plot(pa["Component"][:k], pa["Observed_eigenvalue"][:k], "-o", color=C_SERIES[0], lw=2,
            label="Observed eigenvalues")
    ax.plot(pa["Component"][:k], pa["Random_95th"][:k], "--s", color=C_SERIES[1], lw=2,
            label="Parallel analysis (95th percentile)")
    ax.axhline(1, color=C_TEXT2, lw=1, ls=":")
    ax.set_xlabel("Component")
    ax.set_ylabel("Eigenvalue")
    ax.set_title("Scree plot with parallel analysis")
    ax.legend()
    save(fig, folder, "Fig05_scree_parallel", figs)


def fig_loadings(load_df: pd.DataFrame, folder, figs, name="Fig06_outer_loadings", title="Outer loadings"):
    fig, ax = plt.subplots(figsize=(7, 9))
    y = np.arange(len(load_df))[::-1]
    colors = [C_SERIES[list(CONSTRUCTS).index(c) % 8] if c in CONSTRUCTS else C_SERIES[0]
              for c in load_df["Construct"]]
    ax.barh(y, load_df["Loading"], color=colors, height=0.65)
    ax.axvline(0.708, color=C_TEXT, lw=1, ls="--")
    ax.axvline(0.40, color=C_TEXT2, lw=1, ls=":")
    ax.text(0.708, len(load_df) + 0.2, "0.708", ha="center", fontsize=8, color=C_TEXT)
    ax.text(0.40, len(load_df) + 0.2, "0.40", ha="center", fontsize=8, color=C_TEXT2)
    ax.set_yticks(y)
    ax.set_yticklabels([f"{i} ({c})" for i, c in zip(load_df["Indicator"], load_df["Construct"])],
                       fontsize=8)
    for yy, v in zip(y, load_df["Loading"]):
        ax.text(v + 0.01, yy, f"{v:.3f}", va="center", fontsize=7, color=C_TEXT)
    ax.set_xlim(0, 1.08)
    ax.set_xlabel("Standardised outer loading")
    ax.set_title(title)
    ax.grid(axis="y", visible=False)
    save(fig, folder, name, figs)


def fig_matrix(M: np.ndarray, labels: list[str], title: str, folder, figs, name: str,
               thresholds=(0.85, 0.90), diag: np.ndarray | None = None):
    L = len(labels)
    fig, ax = plt.subplots(figsize=(6.5, 5.5))
    show = np.where(np.isnan(M), np.nan, M)
    im = ax.imshow(show, cmap=_seq_cmap(), vmin=0, vmax=max(1.0, np.nanmax(show)))
    for i in range(L):
        for j in range(L):
            v = M[i, j]
            if np.isfinite(v):
                flag = ""
                if thresholds and i != j:
                    flag = " !!" if v > thresholds[1] else " !" if v > thresholds[0] else ""
                col = "white" if v > 0.6 else C_TEXT
                ax.text(j, i, f"{v:.3f}{flag}", ha="center", va="center", fontsize=8, color=col,
                        fontweight="bold" if (diag is not None and i == j) else "normal")
    ax.set_xticks(range(L))
    ax.set_yticks(range(L))
    ax.set_xticklabels(labels)
    ax.set_yticklabels(labels)
    ax.grid(False)
    fig.colorbar(im, ax=ax, shrink=0.8)
    ax.set_title(title)
    save(fig, folder, name, figs)


def _node_positions(spec: ModelSpec):
    ex = [spec.lv[j] for j in spec.exog]
    en = [spec.lv[j] for j in spec.endog]
    mids = [spec.lv[j] for j in spec.endog if spec.succ[j]]
    finals = [c for c in en if c not in mids]
    pos = {}
    cols = [ex] + ([mids] if mids else []) + [finals]
    xs = np.linspace(0.18, 0.82, len(cols))
    for x, col in zip(xs, cols):
        ys = np.linspace(0.85, 0.15, len(col)) if len(col) > 1 else [0.5]
        for c, y in zip(col, ys):
            pos[c] = (x, y)
    return pos


def fig_structural(spec: ModelSpec, path_tab: pd.DataFrame, r2: dict, folder, figs,
                   name="Fig09_structural_model", title="Structural model"):
    pos = _node_positions(spec)
    fig, ax = plt.subplots(figsize=(10, 7.5))
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    w, h = 0.17, 0.085
    for c, (x, y) in pos.items():
        is_endo = spec.lv_index[c] in spec.endog
        ax.add_patch(Ellipse((x, y), w, h * 1.25, fc="#eaf2fc" if is_endo else "#fdf0ea",
                             ec=C_SERIES[0] if is_endo else C_SERIES[1], lw=1.8, zorder=3))
        full = CONSTRUCTS[c]["name"] if c in CONSTRUCTS else c
        ax.text(x, y + 0.012, c, ha="center", va="center", fontsize=11, fontweight="bold",
                color=C_TEXT, zorder=4)
        sub = f"R² = {r2[c]:.3f}" if c in r2 else full.replace(" ", "\n", 1) if len(full) > 18 else full
        ax.text(x, y - 0.022, sub, ha="center", va="center", fontsize=7.5, color=C_TEXT2, zorder=4)
    for _, row in path_tab.iterrows():
        s, t = row["Source"], row["Target"]
        (x1, y1), (x2, y2) = pos[s], pos[t]
        sig = row["p_value"] < ALPHA
        arr = FancyArrowPatch((x1 + w / 2, y1), (x2 - w / 2, y2), arrowstyle="-|>", mutation_scale=14,
                              color=C_TEXT if sig else "#a8a69f", lw=1.4 + 3 * min(abs(row["Original_sample"]), 1),
                              ls="-" if sig else "--", zorder=2, shrinkA=2, shrinkB=2)
        ax.add_patch(arr)
        frac = 0.35 if y2 >= y1 else 0.62
        tx, ty = x1 + w / 2 + (x2 - x1 - w) * frac, y1 + (y2 - y1) * frac
        lab = f"{row.get('Hypothesis', '')} β={row['Original_sample']:.3f}\np {fmt_p(row['p_value'])}"
        ax.text(tx, ty, lab.strip(), ha="center", va="center", fontsize=7.5, color=C_TEXT,
                bbox=dict(boxstyle="round,pad=0.25", fc="white", ec=C_GRID, lw=0.8), zorder=5)
    ax.set_title(f"{title}\n(solid = significant at p < {ALPHA}; dashed = not significant; "
                 f"line width ∝ |β|)", fontsize=10)
    save(fig, folder, name, figs)


def fig_full_model(spec: ModelSpec, res: PLSResult, path_tab: pd.DataFrame, folder, figs):
    pos = _node_positions(spec)
    fig, ax = plt.subplots(figsize=(15, 11))
    ax.set_xlim(-0.12, 1.12)
    ax.set_ylim(-0.05, 1.05)
    ax.axis("off")
    for j, c in enumerate(spec.lv):
        x, y = pos[c]
        is_endo = j in spec.endog
        ax.add_patch(Ellipse((x, y), 0.11, 0.075, fc="#eaf2fc" if is_endo else "#fdf0ea",
                             ec=C_SERIES[0] if is_endo else C_SERIES[1], lw=1.8, zorder=3))
        ax.text(x, y + 0.008, c, ha="center", fontsize=10, fontweight="bold", zorder=4, color=C_TEXT)
        if is_endo:
            ax.text(x, y - 0.018, f"R²={res.R2[j]:.3f}", ha="center", fontsize=7, zorder=4, color=C_TEXT2)
        b = spec.block_idx[j]
        side = -1 if j in spec.exog else 1
        bx = x + side * 0.16
        ys = np.linspace(y + 0.014 * len(b), y - 0.014 * len(b), len(b)) if len(b) > 1 else [y]
        for idx, yy in zip(b, ys):
            ax.add_patch(FancyBboxPatch((bx - 0.025, yy - 0.009), 0.05, 0.018,
                                        boxstyle="round,pad=0.002", fc="#fffbe6", ec="#c98500", lw=0.8))
            ax.text(bx, yy, spec.indicators[idx], ha="center", va="center", fontsize=6.5, color=C_TEXT)
            ax.annotate("", xy=(bx - side * 0.026, yy), xytext=(x + side * 0.055, y),
                        arrowprops=dict(arrowstyle="-|>", color="#a8a69f", lw=0.6))
            mx = 0.3 * (x + side * 0.055) + 0.7 * (bx - side * 0.026)
            my = 0.3 * y + 0.7 * yy
            ax.text(mx, my, f"{res.loadings[idx]:.2f}", fontsize=5.5, color=C_TEXT2, ha="center",
                    va="center", bbox=dict(fc="white", ec="none", pad=0.3))
    for _, row in path_tab.iterrows():
        (x1, y1), (x2, y2) = pos[row["Source"]], pos[row["Target"]]
        sig = row["p_value"] < ALPHA
        ax.add_patch(FancyArrowPatch((x1 + 0.055, y1), (x2 - 0.055, y2), arrowstyle="-|>",
                                     mutation_scale=13, color=C_TEXT if sig else "#a8a69f",
                                     lw=1.2 + 2.5 * min(abs(row["Original_sample"]), 1),
                                     ls="-" if sig else "--", zorder=2))
        tx, ty = x1 + (x2 - x1) * 0.45, y1 + (y2 - y1) * 0.45
        ax.text(tx, ty, f"{row['Original_sample']:.3f}\n({fmt_p(row['p_value'])})", fontsize=7,
                ha="center", va="center", color=C_TEXT,
                bbox=dict(boxstyle="round,pad=0.2", fc="white", ec=C_GRID), zorder=5)
    ax.set_title("Full PLS path model: outer loadings, path coefficients (p-values) and R²",
                 fontsize=11, fontweight="bold")
    save(fig, folder, "Fig10_full_path_model", figs)


def fig_bootstrap_dist(boot: np.ndarray, path_tab: pd.DataFrame, folder, figs):
    k = len(path_tab)
    cols = 4
    rows = math.ceil(k / cols)
    fig, axes = plt.subplots(rows, cols, figsize=(3.2 * cols, 2.6 * rows))
    axes = np.atleast_1d(axes).ravel()
    for i, (_, r) in enumerate(path_tab.iterrows()):
        ax = axes[i]
        ax.hist(boot[:, i], bins=40, color=C_SERIES[0], edgecolor="white")
        ax.axvline(r["Original_sample"], color=C_TEXT, lw=1.5)
        ax.axvline(r["CI_2.5_percentile"], color=C_SERIES[1], lw=1, ls="--")
        ax.axvline(r["CI_97.5_percentile"], color=C_SERIES[1], lw=1, ls="--")
        ax.axvline(0, color=C_TEXT2, lw=0.8, ls=":")
        ax.set_title(f"{r.get('Hypothesis', '')} {r['Source']} → {r['Target']}", fontsize=9)
        ax.grid(axis="x", visible=False)
    for ax in axes[k:]:
        ax.axis("off")
    fig.suptitle("Bootstrap distributions of path coefficients (dashed = 95% percentile CI)",
                 fontweight="bold", color=C_TEXT)
    fig.tight_layout()
    save(fig, folder, "Fig11_bootstrap_distributions", figs)


def fig_forest(comp: pd.DataFrame, folder, figs):
    models = [c for c in ["Main PLS", "PLSc", "CB-SEM (ML)", "Purified", "Sum-score OLS", "Outliers removed",
                          "With controls"] if f"{c}_beta" in comp.columns]
    fig, ax = plt.subplots(figsize=(8, 6))
    n = len(comp)
    off = np.linspace(-0.3, 0.3, len(models)) if len(models) > 1 else [0]
    for m, (model, o) in enumerate(zip(models, off)):
        y = np.arange(n)[::-1] + o
        b = comp[f"{model}_beta"].values
        lo = comp.get(f"{model}_lo", pd.Series([np.nan] * n)).values
        hi = comp.get(f"{model}_hi", pd.Series([np.nan] * n)).values
        ax.errorbar(b, y, xerr=[b - lo, hi - b] if np.all(np.isfinite(lo)) else None, fmt="o",
                    color=C_SERIES[m], ms=6, elinewidth=1.5, capsize=0, label=model)
    ax.axvline(0, color=C_TEXT2, lw=1)
    ax.set_yticks(np.arange(n)[::-1])
    ax.set_yticklabels(comp["Hypothesis"] + "  " + comp["Path"])
    ax.set_xlabel("Standardised path coefficient (95% CI where available)")
    ax.set_title("Robustness of path estimates across model specifications")
    ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.1), ncol=min(len(models), 6), fontsize=8)
    save(fig, folder, "Fig12_robustness_forest", figs)


def fig_effects(r2_tab: pd.DataFrame, f2_tab: pd.DataFrame, folder, figs):
    fig, axes = plt.subplots(1, 2, figsize=(11, 4))
    ax = axes[0]
    x = np.arange(len(r2_tab))
    ax.bar(x - 0.2, r2_tab["R2"], 0.38, color=C_SERIES[0], label="R²")
    ax.bar(x + 0.2, r2_tab["Q2_blindfolding"], 0.38, color=C_SERIES[1], label="Q² (blindfolding)")
    for xx, a, b in zip(x, r2_tab["R2"], r2_tab["Q2_blindfolding"]):
        ax.text(xx - 0.2, a + 0.01, f"{a:.3f}", ha="center", fontsize=7)
        ax.text(xx + 0.2, b + 0.01, f"{b:.3f}", ha="center", fontsize=7)
    for v in (0.25, 0.50, 0.75):
        ax.axhline(v, color=C_GRID, lw=1, ls="--")
    ax.set_xticks(x)
    ax.set_xticklabels(r2_tab["Construct"])
    ax.set_ylim(0, 1.05)
    ax.set_title("Explanatory power (R²) and predictive relevance (Q²)")
    ax.legend()
    ax.grid(axis="x", visible=False)
    ax = axes[1]
    y = np.arange(len(f2_tab))[::-1]
    ax.barh(y, f2_tab["f2"], color=C_SERIES[0], height=0.6)
    for v, lab in ((0.02, "small"), (0.15, "medium"), (0.35, "large")):
        ax.axvline(v, color=C_TEXT2, lw=0.8, ls="--")
        ax.text(v, len(f2_tab) - 0.3, lab, fontsize=7, color=C_TEXT2, ha="center")
    for yy, v in zip(y, f2_tab["f2"]):
        ax.text(v + 0.01, yy, f"{v:.3f}", va="center", fontsize=7)
    ax.set_yticks(y)
    ax.set_yticklabels(f2_tab["Path"])
    ax.set_title("Effect sizes (f²)")
    ax.grid(axis="y", visible=False)
    fig.tight_layout()
    save(fig, folder, "Fig13_R2_Q2_f2", figs)


def fig_predict(ind: pd.DataFrame, folder, figs):
    fig, ax = plt.subplots(figsize=(9, 4))
    x = np.arange(len(ind))
    ax.bar(x - 0.2, ind["PLS_RMSE"], 0.38, color=C_SERIES[0], label="PLS-SEM RMSE")
    ax.bar(x + 0.2, ind["LM_RMSE"], 0.38, color=C_SERIES[1], label="LM benchmark RMSE")
    ax.set_xticks(x)
    ax.set_xticklabels(ind["Indicator"], rotation=90, fontsize=7)
    ax.set_ylabel("RMSE (original scale)")
    ax.set_title("PLSpredict: out-of-sample prediction errors vs. linear-model benchmark")
    ax.legend()
    ax.grid(axis="x", visible=False)
    save(fig, folder, "Fig14_PLSpredict", figs)


def fig_ipma(ipma: pd.DataFrame, target: str, folder, figs, level="construct"):
    fig, ax = plt.subplots(figsize=(6.5, 5))
    ax.scatter(ipma["Importance"], ipma["Performance"], s=60, color=C_SERIES[0],
               edgecolor="white", linewidth=2, zorder=3)
    for _, r in ipma.iterrows():
        ax.annotate(r["Predictor"], (r["Importance"], r["Performance"]), xytext=(5, 4),
                    textcoords="offset points", fontsize=8, color=C_TEXT)
    ax.axvline(ipma["Importance"].mean(), color=C_TEXT2, lw=1, ls="--")
    ax.axhline(ipma["Performance"].mean(), color=C_TEXT2, lw=1, ls="--")
    ax.set_xlabel("Importance (unstandardised total effect)")
    ax.set_ylabel("Performance (0-100)")
    ax.set_title(f"IPMA ({level} level) - target: {target}")
    save(fig, folder, f"Fig15_IPMA_{level}_{target}", figs)


def fig_fit_boot(dist: np.ndarray, obs: float, hi95: float, hi99: float, folder, figs, kind):
    fig, ax = plt.subplots(figsize=(6, 3.5))
    ax.hist(dist, bins=40, color=C_SERIES[0], edgecolor="white")
    ax.axvline(obs, color=C_TEXT, lw=2, label=f"Observed SRMR = {obs:.3f}")
    ax.axvline(hi95, color=C_SERIES[1], lw=1.5, ls="--", label=f"HI95 = {hi95:.3f}")
    ax.axvline(hi99, color=C_SERIES[7], lw=1.5, ls=":", label=f"HI99 = {hi99:.3f}")
    ax.set_xlabel("SRMR under H0 (Bollen-Stine bootstrap)")
    ax.set_title(f"Exact model fit test - {kind} model")
    ax.legend(fontsize=8)
    ax.grid(axis="x", visible=False)
    save(fig, folder, f"Fig16_fit_bootstrap_{kind}", figs)


def fig_mga(tab: pd.DataFrame, var: str, labels, folder, figs):
    fig, ax = plt.subplots(figsize=(8, 4.5))
    y = np.arange(len(tab))[::-1]
    ax.scatter(tab[f"beta_{labels[0]}"], y + 0.12, color=C_SERIES[0], s=50, label=labels[0], zorder=3)
    ax.scatter(tab[f"beta_{labels[1]}"], y - 0.12, color=C_SERIES[1], s=50, label=labels[1], zorder=3)
    for yy, a, b, p in zip(y, tab[f"beta_{labels[0]}"], tab[f"beta_{labels[1]}"], tab["Permutation_p"]):
        ax.plot([a, b], [yy + 0.12, yy - 0.12], color=C_GRID, lw=1.5, zorder=2)
        ax.text(max(a, b) + 0.03, yy, f"perm. p {fmt_p(p)}", fontsize=7, va="center", color=C_TEXT2)
    ax.set_yticks(y)
    ax.set_yticklabels(tab["Path"])
    ax.axvline(0, color=C_TEXT2, lw=0.8)
    ax.set_xlabel("Path coefficient")
    ax.set_title(f"Multigroup analysis by {var}")
    ax.legend()
    save(fig, folder, f"Fig17_MGA_{var}", figs)


def fig_mahalanobis(md: pd.DataFrame, p: int, folder, figs):
    fig, ax = plt.subplots(figsize=(8, 3.5))
    out = md["Outlier_p<0.001"]
    ax.scatter(md["Case"][~out], md["Mahalanobis_D2"][~out], s=10, color=C_SERIES[0], label="Case")
    ax.scatter(md["Case"][out], md["Mahalanobis_D2"][out], s=24, color=C_SERIES[7],
               label="Multivariate outlier (p < .001)")
    ax.axhline(stats.chi2.ppf(0.999, p), color=C_TEXT2, ls="--", lw=1)
    ax.set_xlabel("Case")
    ax.set_ylabel("Mahalanobis D²")
    ax.set_title("Multivariate outlier screening")
    ax.legend(fontsize=8)
    save(fig, folder, "Fig18_mahalanobis", figs)


# ---- 3D figures -----------------------------------------------------------------

def _style_3d(ax, xl, yl, zl):
    ax.set_xlabel(xl, labelpad=6, color=C_TEXT2)
    ax.set_ylabel(yl, labelpad=6, color=C_TEXT2)
    ax.set_zlabel(zl, labelpad=6, color=C_TEXT2)
    for axis in (ax.xaxis, ax.yaxis, ax.zaxis):
        axis.set_pane_color((0.99, 0.99, 0.985, 1.0))
        axis._axinfo["grid"]["color"] = C_GRID
        axis._axinfo["grid"]["linewidth"] = 0.5
    ax.tick_params(labelsize=7, colors=C_TEXT2)


def _quad_design(a, b):
    return np.column_stack([np.ones_like(a), a, b, a ** 2, b ** 2, a * b])


def fig3d_response_surfaces(res: PLSResult, folder, figs):
    """Quadratic response surface of each outcome over the two exogenous constructs."""
    spec = res.spec
    if len(spec.exog) < 2:
        return
    i1, i2 = spec.exog[:2]
    x, y = res.Y[:, i1], res.Y[:, i2]
    g = np.linspace(-2.6, 2.6, 45)
    GX, GY = np.meshgrid(g, g)
    fig = plt.figure(figsize=(13, 10.5))
    for k, j in enumerate(spec.endog[:4]):
        ax = fig.add_subplot(2, 2, k + 1, projection="3d")
        z = res.Y[:, j]
        coef = np.linalg.lstsq(_quad_design(x, y), z, rcond=None)[0]
        r2 = 1 - np.sum((z - _quad_design(x, y) @ coef) ** 2) / np.sum((z - z.mean()) ** 2)
        GZ = (_quad_design(GX.ravel(), GY.ravel()) @ coef).reshape(GX.shape)
        surf = ax.plot_surface(GX, GY, GZ, cmap=_seq_cmap(), alpha=0.88, linewidth=0, antialiased=True,
                               rstride=1, cstride=1)
        ax.contour(GX, GY, GZ, zdir="z", offset=GZ.min() - 1.0, cmap=_seq_cmap(), levels=12, linewidths=1)
        ax.scatter(x, y, z, s=9, color=C_SERIES[1], alpha=0.75, depthshade=True, edgecolor="white", linewidth=0.3)
        ax.set_zlim(GZ.min() - 1.0, max(GZ.max(), z.max()) + 0.2)
        _style_3d(ax, spec.lv[i1], spec.lv[i2], spec.lv[j])
        ax.view_init(elev=24, azim=-132)
        ax.set_title(f"{CONSTRUCTS.get(spec.lv[j], {'name': spec.lv[j]})['name']}\n"
                     f"quadratic response surface, R² = {r2:.3f}", fontsize=10)
        fig.colorbar(surf, ax=ax, shrink=0.55, pad=0.08, label=f"Predicted {spec.lv[j]} (z-score)")
    fig.suptitle("3D response surfaces: accountability dimensions as a function of control activities "
                 "and financial management\n(points = latent variable scores; floor = contour projection)",
                 fontweight="bold", color=C_TEXT)
    fig.tight_layout()
    save(fig, folder, "Fig20_3D_response_surfaces", figs)


def fig3d_path_bars(path_tab: pd.DataFrame, folder, figs):
    src = list(OrderedDict.fromkeys(path_tab["Source"]))
    tgt = list(OrderedDict.fromkeys(path_tab["Target"]))
    fig = plt.figure(figsize=(9, 7))
    ax = fig.add_subplot(111, projection="3d")
    cmap = _seq_cmap()
    xs, ys, hs, cols = [], [], [], []
    bmax = max(path_tab["Original_sample"].abs().max(), 1e-6)
    for _, r in path_tab.iterrows():
        xs.append(src.index(r["Source"]) - 0.22)
        ys.append(tgt.index(r["Target"]) - 0.22)
        hs.append(r["Original_sample"])
        sig = r["p_value"] < ALPHA
        cols.append(cmap(0.35 + 0.65 * abs(r["Original_sample"]) / bmax) if sig else "#c3c2b7")
    ax.bar3d(xs, ys, np.zeros(len(xs)), 0.44, 0.44, hs, color=cols, alpha=0.95, edgecolor="white",
             linewidth=0.5, shade=True)
    for x_, y_, h_, (_, r) in zip(xs, ys, hs, path_tab.iterrows()):
        ax.text(x_ + 0.22, y_ + 0.22, h_ + 0.06,
                f"{h_:.3f}{stars(r['p_value']) if r['p_value'] < ALPHA else ''}",
                ha="center", fontsize=8, color=C_TEXT, zorder=10)
    ax.set_xticks(range(len(src)))
    ax.set_xticklabels(src)
    ax.set_yticks(range(len(tgt)))
    ax.set_yticklabels(tgt)
    _style_3d(ax, "Predictor", "Outcome", "Path coefficient (β)")
    ax.view_init(elev=28, azim=-55)
    ax.set_title("3D map of standardised path coefficients\n(blue shade ∝ |β| for significant paths; grey = not significant)")
    save(fig, folder, "Fig21_3D_path_coefficients", figs)


def fig3d_htmt(H: np.ndarray, labels: list[str], folder, figs):
    L = len(labels)
    fig = plt.figure(figsize=(9, 7.5))
    ax = fig.add_subplot(111, projection="3d")
    xs, ys, hs, cols = [], [], [], []
    for i in range(L):
        for j in range(i):
            v = H[i, j]
            if not np.isfinite(v):
                continue
            xs.append(j - 0.3)
            ys.append(i - 0.3)
            hs.append(v)
            cols.append(C_SERIES[7] if v > 0.90 else C_SERIES[3] if v > 0.85 else C_SERIES[0])
    ax.bar3d(xs, ys, np.zeros(len(xs)), 0.6, 0.6, hs, color=cols, alpha=0.95, edgecolor="white",
             linewidth=0.5, shade=True)
    for x_, y_, h_ in zip(xs, ys, hs):
        ax.text(x_ + 0.3, y_ + 0.3, h_ + 0.04, f"{h_:.2f}", ha="center", fontsize=7, color=C_TEXT, zorder=10)
    edge = [-0.5, L - 0.5]
    for level, col, ls in ((0.85, C_SERIES[3], "--"), (0.90, C_SERIES[7], "-")):
        xx = [edge[0], edge[1], edge[1], edge[0], edge[0]]
        yy = [edge[0], edge[0], edge[1], edge[1], edge[0]]
        ax.plot(xx, yy, [level] * 5, color=col, lw=1.6, ls=ls, label=f"HTMT = {level:.2f}")
    ax.legend(loc="upper left", fontsize=8)
    ax.set_xticks(range(L))
    ax.set_xticklabels(labels, fontsize=8)
    ax.set_yticks(range(L))
    ax.set_yticklabels(labels, fontsize=8)
    _style_3d(ax, "", "", "HTMT")
    ax.set_zlim(0, max(1.05, np.nanmax(H) + 0.1))
    ax.view_init(elev=26, azim=-60)
    ax.set_title("3D HTMT matrix with 0.85 and 0.90 threshold frames\n"
                 "bars: blue < 0.85, amber 0.85-0.90, red > 0.90")
    save(fig, folder, "Fig22_3D_HTMT", figs)


def fig3d_scatter(res: PLSResult, folder, figs):
    spec = res.spec
    if len(spec.exog) < 2:
        return
    i1, i2 = spec.exog[:2]
    acc = res.Y[:, spec.endog].mean(axis=1)
    fig = plt.figure(figsize=(9, 7.5))
    ax = fig.add_subplot(111, projection="3d")
    sc = ax.scatter(res.Y[:, i1], res.Y[:, i2], acc, c=acc, cmap=_seq_cmap(), s=28, edgecolor="white",
                    linewidth=0.4, depthshade=True)
    A = np.column_stack([np.ones(len(acc)), res.Y[:, i1], res.Y[:, i2]])
    b = np.linalg.lstsq(A, acc, rcond=None)[0]
    g = np.linspace(-2.6, 2.6, 20)
    GX, GY = np.meshgrid(g, g)
    ax.plot_surface(GX, GY, b[0] + b[1] * GX + b[2] * GY, color=C_SERIES[0], alpha=0.18)
    _style_3d(ax, spec.lv[i1], spec.lv[i2], "Overall accountability (mean of LV scores)")
    ax.view_init(elev=20, azim=-125)
    fig.colorbar(sc, ax=ax, shrink=0.6, pad=0.1, label="Overall accountability")
    ax.set_title("3D scatter of latent variable scores with the fitted regression plane")
    save(fig, folder, "Fig23_3D_LV_scatter", figs)


def fig3d_cross_loadings(cl: np.ndarray, spec: ModelSpec, folder, figs):
    fig = plt.figure(figsize=(12, 7))
    ax = fig.add_subplot(111, projection="3d")
    k, L = cl.shape
    xs, ys, hs, cols = [], [], [], []
    for i in range(k):
        for j in range(L):
            own = spec.ind_lv[i] == j
            xs.append(i - 0.4)
            ys.append(j - 0.4)
            hs.append(abs(cl[i, j]))
            cols.append(C_SERIES[list(CONSTRUCTS).index(spec.lv[j]) % 8] if own and spec.lv[j] in CONSTRUCTS
                        else C_SERIES[0] if own else "#dddcd7")
    ax.bar3d(xs, ys, np.zeros(len(xs)), 0.8, 0.8, hs, color=cols, alpha=0.95, edgecolor="white",
             linewidth=0.3, shade=True)
    ax.set_xticks(range(k))
    ax.set_xticklabels(spec.indicators, rotation=90, fontsize=6)
    ax.set_yticks(range(L))
    ax.set_yticklabels(spec.lv, fontsize=8)
    _style_3d(ax, "", "", "|Loading|")
    ax.view_init(elev=32, azim=-70)
    ax.set_title("3D cross-loading landscape (coloured = own construct, grey = other constructs)")
    save(fig, folder, "Fig24_3D_cross_loadings", figs)


def interactive_3d(res: PLSResult, path_tab: pd.DataFrame, H: np.ndarray, path: Path) -> bool:
    """Rotatable 3D figures in one HTML file (requires plotly)."""
    try:
        import plotly.graph_objects as go
        from plotly.subplots import make_subplots
    except ImportError:
        return False
    spec = res.spec
    i1, i2 = spec.exog[:2]
    x, y = res.Y[:, i1], res.Y[:, i2]
    g = np.linspace(-2.6, 2.6, 40)
    GX, GY = np.meshgrid(g, g)
    titles = [f"{spec.lv[j]} response surface" for j in spec.endog[:4]]
    fig = make_subplots(rows=2, cols=2, specs=[[{"type": "scene"}] * 2] * 2, subplot_titles=titles,
                        horizontal_spacing=0.02, vertical_spacing=0.06)
    scale = [[0, C_BLUE_SEQ[0]], [0.5, C_BLUE_SEQ[3]], [1, C_BLUE_SEQ[-1]]]
    for k, j in enumerate(spec.endog[:4]):
        z = res.Y[:, j]
        coef = np.linalg.lstsq(_quad_design(x, y), z, rcond=None)[0]
        GZ = (_quad_design(GX.ravel(), GY.ravel()) @ coef).reshape(GX.shape)
        r, c = divmod(k, 2)
        fig.add_trace(go.Surface(x=g, y=g, z=GZ, colorscale=scale, opacity=0.9, showscale=False,
                                 name=spec.lv[j]), row=r + 1, col=c + 1)
        fig.add_trace(go.Scatter3d(x=x, y=y, z=z, mode="markers", name=f"{spec.lv[j]} scores",
                                   marker=dict(size=3, color=C_SERIES[1], line=dict(width=0.5, color="white")),
                                   hovertemplate=f"{spec.lv[i1]}=%{{x:.2f}}<br>{spec.lv[i2]}=%{{y:.2f}}<br>"
                                                 f"{spec.lv[j]}=%{{z:.2f}}<extra></extra>"),
                      row=r + 1, col=c + 1)
        fig.update_scenes(dict(xaxis_title=spec.lv[i1], yaxis_title=spec.lv[i2], zaxis_title=spec.lv[j]),
                          row=r + 1, col=c + 1)
    fig.update_layout(title="Interactive 3D response surfaces (drag to rotate, scroll to zoom)",
                      height=950, showlegend=False, paper_bgcolor="white", font=dict(color=C_TEXT))
    fig2 = go.Figure()
    src = list(OrderedDict.fromkeys(path_tab["Source"]))
    tgt = list(OrderedDict.fromkeys(path_tab["Target"]))
    for _, rr in path_tab.iterrows():
        xi, yi, h = src.index(rr["Source"]), tgt.index(rr["Target"]), rr["Original_sample"]
        col = C_SERIES[0] if rr["p_value"] < ALPHA else "#c3c2b7"
        xs = [xi - .3, xi + .3, xi + .3, xi - .3, xi - .3, xi + .3, xi + .3, xi - .3]
        ys = [yi - .3, yi - .3, yi + .3, yi + .3, yi - .3, yi - .3, yi + .3, yi + .3]
        zs = [0, 0, 0, 0, h, h, h, h]
        fig2.add_trace(go.Mesh3d(x=xs, y=ys, z=zs, i=[7, 0, 0, 0, 4, 4, 6, 6, 4, 0, 3, 2],
                                 j=[3, 4, 1, 2, 5, 6, 5, 2, 0, 1, 6, 3], k=[0, 7, 2, 3, 6, 7, 1, 1, 5, 5, 7, 6],
                                 color=col, opacity=0.95, flatshading=True,
                                 hovertext=f"{rr['Source']} → {rr['Target']}: β = {h:.3f}, p = {fmt_p(rr['p_value'])}",
                                 hoverinfo="text", name=f"{rr['Source']}→{rr['Target']}"))
    fig2.update_layout(title="Interactive 3D path coefficients (blue = significant)", height=650,
                       scene=dict(xaxis=dict(tickvals=list(range(len(src))), ticktext=src, title="Predictor"),
                                  yaxis=dict(tickvals=list(range(len(tgt))), ticktext=tgt, title="Outcome"),
                                  zaxis=dict(title="β")), showlegend=False)
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("<html><head><meta charset='utf-8'><title>Interactive 3D PLS-SEM figures</title></head><body>")
        fh.write(fig.to_html(full_html=False, include_plotlyjs=True))
        fh.write(fig2.to_html(full_html=False, include_plotlyjs=False))
        fh.write("</body></html>")
    return True


# =============================================================================
# 9. REPORT WRITERS
# =============================================================================

def write_excel(tables: "OrderedDict[str, tuple[pd.DataFrame, str]]", path: Path) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    used = set()
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        index_rows = []
        for name, (df, desc) in tables.items():
            sheet = name[:31]
            base, k = sheet, 1
            while sheet in used:
                sheet = f"{base[:28]}_{k}"
                k += 1
            used.add(sheet)
            index_rows.append([sheet, desc])
            df.to_excel(xw, sheet_name=sheet, index=False, startrow=2)
            ws = xw.sheets[sheet]
            ws["A1"] = name
            ws["A1"].font = Font(bold=True, size=12)
            ws["A2"] = desc
            ws["A2"].font = Font(italic=True, color="52514E")
            for c in ws[3]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="256ABF")
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for col_idx, col in enumerate(df.columns, 1):
                width = max(len(str(col)), *(len(f"{v:.3f}" if isinstance(v, float) else str(v))
                                             for v in df[col].head(200))) if len(df) else len(str(col))
                ws.column_dimensions[get_column_letter(col_idx)].width = min(max(10, width + 2), 60)
                for row in ws.iter_rows(min_row=4, min_col=col_idx, max_col=col_idx):
                    for c in row:
                        if isinstance(c.value, float):
                            c.number_format = "0.000"
            ws.freeze_panes = "B4"
        idx = pd.DataFrame(index_rows, columns=["Sheet", "Contents"])
        idx.to_excel(xw, sheet_name="INDEX", index=False)
        ws = xw.sheets["INDEX"]
        ws.column_dimensions["A"].width = 34
        ws.column_dimensions["B"].width = 120
        xw.book.move_sheet("INDEX", offset=-len(xw.book.sheetnames) + 1)


def write_docx(doc_tables: list, figs: list, summary_lines: list, path: Path) -> bool:
    try:
        import docx
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt
    except ImportError:
        return False
    d = docx.Document()
    style = d.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(10)
    d.add_heading("PLS-SEM Results Report", 0)
    d.add_heading("Key diagnostics", 1)
    for line in summary_lines:
        d.add_paragraph(line, style="List Bullet")
    for title, df in doc_tables:
        d.add_heading(title, 2)
        t = d.add_table(rows=1, cols=len(df.columns))
        t.style = "Light Grid Accent 1"
        for i, c in enumerate(df.columns):
            t.rows[0].cells[i].text = str(c)
        for _, r in df.iterrows():
            cells = t.add_row().cells
            for i, v in enumerate(r):
                if isinstance(v, (float, np.floating)):
                    cells[i].text = "" if not np.isfinite(v) else f"{v:.3f}"
                else:
                    cells[i].text = str(v)
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(8)
    d.add_page_break()
    d.add_heading("Figures", 1)
    for f in figs:
        d.add_paragraph(f.stem.replace("_", " "))
        d.add_picture(str(f), width=Inches(6.2))
        d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.save(path)
    return True


# =============================================================================
# 10. MAIN ANALYSIS
# =============================================================================

def main() -> None:
    t_start = time.time()
    setup_matplotlib()
    in_path = locate_input()
    out_dir = in_path.parent / OUTPUT_FOLDER_NAME
    fig_dir = out_dir / "figures"
    fig_dir.mkdir(parents=True, exist_ok=True)
    log(f"Input : {in_path}")
    log(f"Output: {out_dir}")

    tables: "OrderedDict[str, tuple[pd.DataFrame, str]]" = OrderedDict()
    figs: list[Path] = []
    summary: list[str] = []

    def add(name, df, desc):
        tables[name] = (df.reset_index(drop=True) if isinstance(df, pd.DataFrame) else df, desc)

    # ---------------- data ----------------
    raw, items_all, demo, screening = load_data(in_path)
    blocks_full = OrderedDict((c, list(v["items"])) for c, v in CONSTRUCTS.items())
    blocks = OrderedDict((c, [i for i in v if i not in DROP_ITEMS]) for c, v in blocks_full.items())
    paths = [(s, t) for s, t, _ in HYPOTHESES]
    hyp_label = {(s, t): h for s, t, h in HYPOTHESES}
    spec = ModelSpec(blocks, paths)
    items = items_all[spec.indicators]
    X = items.values
    n, p = X.shape
    add("S1_Missing_values", screening["missing"], "Missing and out-of-range values per item "
        f"(strategy: {MISSING_STRATEGY}).")
    add("S2_Straight_lining", screening["straight_lining"],
        "Row-wise SD across all items; SD = 0 indicates straight-lining.")
    md = mahalanobis(items)
    add("S3_Mahalanobis", md, "Mahalanobis D² for each case; p < .001 flags multivariate outliers.")
    fig_mahalanobis(md, p, fig_dir, figs)
    n_out = int(md["Outlier_p<0.001"].sum())
    summary.append(f"Sample: n = {n}, indicators = {p}, multivariate outliers (p<.001) = {n_out}, "
                   f"straight-liners = {int(screening['straight_lining']['Straight_liner'].sum())}.")

    # ---------------- descriptives ----------------
    log("Descriptive statistics and normality")
    demo_tab = demographic_table(demo)
    add("T1_Demographics", demo_tab, "Table 1. Demographic profile of the respondents.")
    fig_demographics(demo_tab, fig_dir, figs)
    desc = descriptive_items(items_all[[i for i in spec.indicators]])
    add("T2_Item_descriptives", desc, "Item descriptives and univariate normality "
        "(skewness, excess kurtosis, Shapiro-Wilk, Kolmogorov-Smirnov).")
    means = construct_means(items, blocks)
    cdesc = descriptive_constructs(means)
    add("T2b_Construct_descriptives", cdesc, "Construct mean-score descriptives and normality.")
    mard = mardia(X)
    add("T2c_Mardia", mard, "Mardia's multivariate skewness and kurtosis.")
    fig_item_means(desc, fig_dir, figs)
    fig_normality(means, fig_dir, figs)
    fig_item_corr(items, fig_dir, figs)
    non_norm = int((desc["Shapiro_p"] < 0.05).sum())
    summary.append(f"Normality: {non_norm}/{p} items deviate from normality (Shapiro-Wilk p<.05); "
                   f"Mardia skewness p = {fmt_p(mard.loc[0, 'p_value'])}, kurtosis p = "
                   f"{fmt_p(mard.loc[1, 'p_value'])} -> supports a non-parametric (PLS-SEM) approach.")

    # ---------------- dimensionality / CMB ----------------
    log("Dimensionality: KMO, Bartlett, parallel analysis, EFA, Harman")
    kmo_sum, msa = kmo_bartlett(X, spec.indicators)
    add("D1_KMO_Bartlett", kmo_sum, "Sampling adequacy (KMO) and Bartlett's test of sphericity.")
    add("D2_Item_MSA", msa, "Measure of sampling adequacy per item (>= 0.50 acceptable).")
    pa, k_pa = parallel_analysis(X, N_PARALLEL, SEED)
    add("D3_Parallel_analysis", pa, f"Eigenvalues vs. random data (Horn). Factors retained: {k_pa}.")
    fig_scree(pa, fig_dir, figs)
    harman = pa.loc[0, "Variance_%"]
    add("D4_Harman_single_factor", pd.DataFrame([[pa.loc[0, "Observed_eigenvalue"], harman,
                                                  "Concern" if harman > 50 else "No concern"]],
                                                columns=["First_eigenvalue", "Variance_explained_%",
                                                         "Assessment_(>50% = concern)"]),
        "Harman's single-factor test (unrotated first component).")
    for k_fac, tag in ((len(CONSTRUCTS), "theoretical"), (max(k_pa, 1), "parallel_analysis")):
        add(f"D5_EFA_{tag}_promax", efa_table(X, spec.indicators, k_fac, "promax"),
            f"Exploratory factor analysis (principal axis, promax) with {k_fac} factors ({tag}).")
        add(f"D6_EFA_{tag}_varimax", efa_table(X, spec.indicators, k_fac, "varimax"),
            f"Exploratory factor analysis (principal axis, varimax) with {k_fac} factors ({tag}).")

    # ---------------- main PLS ----------------
    log("Estimating main PLS model")
    res = fit_pls(X, spec)
    if not res.converged:
        log("WARNING: PLS algorithm did not converge")
    R = corr(X)
    path_ids = [(spec.lv_index[s], spec.lv_index[t]) for s, t in spec.paths]
    L = spec.L

    def extract_main(r: PLSResult):
        pathv = [r.B[i, j] for i, j in path_ids]
        tot = r.total
        totv = [tot[i, j] for i in range(L) for j in range(L) if i != j and tot[i, j] != 0]
        ind = r.indirect
        r2 = [r.R2[j] for j in spec.endog]
        f2 = f2_values(r)
        f2v = [f2[(i, j)] for i, j in path_ids]
        H = htmt_matrix((r.Z.T @ r.Z) / r.Z.shape[0], spec)
        hv = H[np.tril_indices(L, -1)]
        return np.concatenate([pathv, r.loadings, r.weights, r2, f2v, hv,
                               [ind[i, j] for i in range(L) for j in range(L) if i != j],
                               totv if len(totv) else []])

    est = extract_main(res)
    log(f"Bootstrapping main model ({N_BOOT} subsamples)")
    boot = bootstrap(X, spec, N_BOOT, SEED, extract_main, label="main")
    log("Jackknife for BCa intervals and case influence")
    jack = jackknife(X, spec, extract_main)
    ci = ci_table(est, boot, jack, n)
    k_p, k_i = len(path_ids), len(spec.indicators)
    sl = OrderedDict()
    pos = 0
    for key, size in (("paths", k_p), ("load", k_i), ("wts", k_i), ("r2", len(spec.endog)),
                      ("f2", k_p), ("htmt", L * (L - 1) // 2), ("ind", L * (L - 1))):
        sl[key] = slice(pos, pos + size)
        pos += size
    sl["tot"] = slice(pos, len(est))

    # ---- measurement model
    rel = reliability(res, R)
    rel["Construct_name"] = [CONSTRUCTS[c]["name"] for c in rel["Construct"]]
    rel["alpha>=0.70"] = rel["Cronbach_alpha"] >= 0.70
    rel["rho_A>=0.70"] = rel["rho_A"] >= 0.70
    rel["rho_C_0.70-0.95"] = (rel["Composite_reliability_rho_C"] >= 0.70) & \
                             (rel["Composite_reliability_rho_C"] <= 0.95)
    rel["AVE>=0.50"] = rel["AVE"] >= 0.50
    ovif = outer_vif(R, spec)
    load_ci = ci.iloc[sl["load"]].reset_index(drop=True)
    wts_ci = ci.iloc[sl["wts"]].reset_index(drop=True)
    load_df = pd.DataFrame({
        "Construct": [spec.lv[j] for j in spec.ind_lv], "Indicator": spec.indicators,
        "Loading": res.loadings, "Indicator_reliability_(loading²)": res.loadings ** 2,
        "Loading_t": load_ci["t_statistic"], "Loading_p": load_ci["p_value"],
        "Loading_CI_2.5": load_ci["CI_2.5_percentile"], "Loading_CI_97.5": load_ci["CI_97.5_percentile"],
        "Outer_weight": res.weights, "Weight_t": wts_ci["t_statistic"], "Weight_p": wts_ci["p_value"],
        "Outer_VIF": ovif})
    load_df["Assessment"] = np.where(load_df["Loading"] >= 0.708, "retain",
                                     np.where(load_df["Loading"] >= 0.40, "consider removal (0.40-0.708)",
                                              "remove (< 0.40)"))
    add("T3_Outer_loadings", load_df, "Table 3. Outer loadings, indicator reliability, bootstrap "
        "significance, outer weights and collinearity (VIF < 5; ideally < 3.3).")
    add("T4_Reliability_validity", rel, "Table 4. Construct reliability and convergent validity "
        "(Cronbach's alpha, rho_A, rho_C, AVE).")
    fig_loadings(load_df, fig_dir, figs)

    cl = cross_loadings(res)
    cl_df = pd.DataFrame(cl, columns=spec.lv)
    cl_df.insert(0, "Indicator", spec.indicators)
    cl_df.insert(1, "Own_construct", [spec.lv[j] for j in spec.ind_lv])
    own = cl[np.arange(k_i), spec.ind_lv]
    other = np.where(np.eye(L, dtype=bool)[spec.ind_lv], -np.inf, cl)
    cl_df["Max_other_loading"] = other.max(axis=1)
    cl_df["Max_other_construct"] = [spec.lv[j] for j in other.argmax(axis=1)]
    cl_df["Own_minus_max_other"] = own - other.max(axis=1)
    cl_df["Problem_(own<=other+0.10)"] = cl_df["Own_minus_max_other"] <= 0.10
    add("T5a_Cross_loadings", cl_df, "Cross-loadings: each indicator should load highest on its own "
        "construct (gap > 0.10 recommended).")

    H = htmt_matrix(R, spec)
    H2 = htmt_matrix(R, spec, "htmt2")
    htmt_df = pd.DataFrame(H, columns=spec.lv, index=spec.lv).reset_index().rename(columns={"index": "Construct"})
    add("T5_HTMT", htmt_df, "Table 5. Heterotrait-monotrait ratio (HTMT). Thresholds 0.85 "
        "(conservative) / 0.90 (liberal).")
    add("T5b_HTMT2", pd.DataFrame(H2, columns=spec.lv, index=spec.lv).reset_index()
        .rename(columns={"index": "Construct"}), "HTMT2 (geometric-mean version, Roemer et al. 2021).")
    hci = ci.iloc[sl["htmt"]].reset_index(drop=True)
    tri = np.tril_indices(L, -1)
    htmt_inf = pd.DataFrame({
        "Pair": [f"{spec.lv[i]} <-> {spec.lv[j]}" for i, j in zip(*tri)],
        "HTMT": H[tri], "Boot_mean": hci["Sample_mean"],
        "CI_5%": np.nanpercentile(boot[:, sl["htmt"]], 5, axis=0),
        "CI_95%_(one-sided_upper)": np.nanpercentile(boot[:, sl["htmt"]], 95, axis=0)})
    htmt_inf["HTMT<0.85"] = htmt_inf["HTMT"] < 0.85
    htmt_inf["HTMT<0.90"] = htmt_inf["HTMT"] < 0.90
    htmt_inf["Upper_CI<0.90"] = htmt_inf["CI_95%_(one-sided_upper)"] < 0.90
    htmt_inf["Upper_CI<1"] = htmt_inf["CI_95%_(one-sided_upper)"] < 1.0
    add("T5c_HTMT_inference", htmt_inf, "Bootstrap HTMT inference (Henseler et al., 2015; "
        "Franke & Sarstedt, 2019): upper bound of the one-sided 95% CI.")
    fig_matrix(H, spec.lv, "HTMT matrix ( ! > 0.85, !! > 0.90 )", fig_dir, figs, "Fig07_HTMT")

    FL = res.C.copy()
    FL[np.triu_indices(L, 1)] = np.nan
    np.fill_diagonal(FL, rel["sqrt_AVE"].values)
    fl_df = pd.DataFrame(FL, columns=spec.lv, index=spec.lv).reset_index().rename(columns={"index": "Construct"})
    fl_ok = []
    for j in range(L):
        others = np.abs(np.delete(res.C[j], j))
        fl_ok.append(rel["sqrt_AVE"].iloc[j] > others.max())
    fl_df["Fornell_Larcker_met"] = fl_ok
    add("T6_Fornell_Larcker", fl_df, "Table 6. Fornell-Larcker criterion (diagonal = sqrt AVE; "
        "off-diagonal = construct correlations).")
    fig_matrix(FL, spec.lv, "Fornell-Larcker (diagonal = √AVE)", fig_dir, figs, "Fig08_Fornell_Larcker",
               thresholds=None, diag=rel["sqrt_AVE"].values)

    lv_corr = pd.DataFrame(res.C, columns=spec.lv, index=spec.lv).reset_index().rename(columns={"index": "Construct"})
    add("T6b_LV_correlations", lv_corr, "Latent variable correlation matrix.")

    # ---- structural model
    log("Structural model assessment")
    pci = ci.iloc[sl["paths"]].reset_index(drop=True)
    path_tab = pd.concat([pd.DataFrame({
        "Hypothesis": [hyp_label.get(pp, "") for pp in spec.paths],
        "Source": [s for s, _ in spec.paths], "Target": [t for _, t in spec.paths],
        "Path": [f"{s} -> {t}" for s, t in spec.paths]}), pci], axis=1)
    path_tab["Significance"] = path_tab["p_value"].apply(stars)
    path_tab["Decision"] = np.where((path_tab["p_value"] < ALPHA) & (path_tab["Original_sample"] > 0),
                                    "Supported", "Not supported")
    f2 = f2_values(res)
    ivif = inner_vif(res)
    path_tab["f2"] = [f2[ij] for ij in path_ids]
    path_tab["f2_effect"] = path_tab["f2"].apply(effect_label_f2)
    path_tab["Inner_VIF"] = [ivif[ij] for ij in path_ids]
    # post-hoc power
    power = []
    for (i, j), f in zip(path_ids, path_tab["f2"]):
        kpred = len(spec.pred[j])
        df1, df2 = 1, n - kpred - 1
        fcrit = stats.f.ppf(1 - ALPHA, df1, df2)
        power.append(float(stats.ncf.sf(fcrit, df1, df2, max(f, 0) * n)))
    path_tab["Post_hoc_power"] = power
    add("T9_Hypotheses_testing", path_tab, f"Table 9. Path coefficients, bootstrapping ({N_BOOT} "
        "subsamples; percentile, bias-corrected and BCa 95% CIs), f², inner VIF and post-hoc power.")
    fig_bootstrap_dist(boot[:, sl["paths"]], path_tab, fig_dir, figs)

    # Q2 blindfolding and q2
    log("Blindfolding (Q²) and q² effect sizes")
    Zfull, _, _ = standardize(X)
    q2 = blindfolding(Zfull, spec)
    q2_rows = []
    for (s, t), (i, j) in zip(spec.paths, path_ids):
        q2_ex = blindfolding(Zfull, spec.without_path(s, t), targets=[j])[j]
        q2_rows.append([hyp_label[(s, t)], f"{s} -> {t}", q2[j], q2_ex, (q2[j] - q2_ex) / (1 - q2[j])])
    q2_eff = pd.DataFrame(q2_rows, columns=["Hypothesis", "Path", "Q2_included", "Q2_excluded", "q2"])
    q2_eff["q2_effect"] = q2_eff["q2"].apply(lambda v: effect_label_f2(v))
    r2_ci = ci.iloc[sl["r2"]].reset_index(drop=True)
    r2_tab = pd.DataFrame({
        "Construct": [spec.lv[j] for j in spec.endog],
        "Name": [CONSTRUCTS[spec.lv[j]]["name"] for j in spec.endog],
        "R2": [res.R2[j] for j in spec.endog], "R2_adjusted": [res.adj_r2()[j] for j in spec.endog],
        "R2_t": r2_ci["t_statistic"], "R2_p": r2_ci["p_value"],
        "R2_CI_2.5": r2_ci["CI_2.5_percentile"], "R2_CI_97.5": r2_ci["CI_97.5_percentile"],
        "Q2_blindfolding": [q2[j] for j in spec.endog]})
    r2_tab["R2_level_(Hair)"] = pd.cut(r2_tab["R2"], [-1, 0.25, 0.50, 0.75, 1.01],
                                       labels=["very weak", "weak", "moderate", "substantial"]).astype(str)
    r2_tab["Q2_relevance"] = r2_tab["Q2_blindfolding"].apply(effect_label_q2)
    add("T10_R2_Q2", r2_tab, "Explanatory power (R², adjusted R², bootstrap CI) and predictive relevance "
        f"(Stone-Geisser Q², blindfolding D = {BLINDFOLD_D}).")
    add("T10b_q2_effect_sizes", q2_eff, "q² effect sizes (change in Q² when a path is omitted).")
    fig_effects(r2_tab, path_tab[["Path", "f2"]], fig_dir, figs)

    # total and indirect effects
    ind_ci = ci.iloc[sl["ind"]].reset_index(drop=True)
    pairs = [(i, j) for i in range(L) for j in range(L) if i != j]
    ind_df = pd.concat([pd.DataFrame({"Effect": [f"{spec.lv[i]} -> {spec.lv[j]}" for i, j in pairs]}),
                        ind_ci], axis=1)
    ind_df = ind_df[ind_df["Original_sample"].abs() > 1e-12]
    tot = res.total
    tot_df = pd.DataFrame([[f"{spec.lv[i]} -> {spec.lv[j]}", tot[i, j]] for i, j in pairs if tot[i, j] != 0],
                          columns=["Effect", "Total_effect"])
    if len(tot_df) and sl["tot"].stop > sl["tot"].start:
        tci = ci.iloc[sl["tot"]].reset_index(drop=True)
        tot_df = pd.concat([tot_df, tci.drop(columns=["Original_sample"])], axis=1)
    add("T11_Total_effects", tot_df, "Total effects with bootstrap inference.")
    if len(ind_df):
        add("T11b_Indirect_effects", ind_df, "Specific indirect effects.")

    # minimum sample size
    bmin = path_tab["Original_sample"].abs().min()
    nmin = math.ceil((2.486 / bmin) ** 2) if bmin > 0 else np.nan
    add("T12_Sample_size_power", pd.DataFrame([[n, bmin, nmin, "Adequate" if n >= nmin else "Insufficient"]],
                                              columns=["Sample_size", "Smallest_|beta|",
                                                       "Minimum_n_inverse_square_root_(5%,80%)",
                                                       "Assessment"]),
        "Minimum sample size: inverse square root method (Kock & Hadaya, 2018).")

    # PLSpredict
    log("PLSpredict (k-fold cross-validation)")
    pred_ind, pred_lv = pls_predict(X, spec)
    add("T13_PLSpredict_indicators", pred_ind, f"PLSpredict ({PREDICT_FOLDS}-fold, {PREDICT_REPEATS} "
        "repetitions): Q²_predict and RMSE/MAE vs. linear-model benchmark.")
    add("T13b_PLSpredict_constructs", pred_lv, "PLSpredict construct-level Q²_predict.")
    share = (pred_ind["RMSE_diff_PLS_minus_LM"] < 0).mean()
    pp = "high" if share == 1 else "medium" if share >= 0.5 else "low" if share > 0 else "none"
    summary.append(f"PLSpredict: Q²_predict > 0 for {(pred_ind['Q2_predict'] > 0).sum()}/{len(pred_ind)} "
                   f"indicators; PLS RMSE < LM RMSE for {share * 100:.0f}% -> {pp} predictive power.")
    fig_predict(pred_ind, fig_dir, figs)

    # ---- model fit
    log("Model fit and Bollen-Stine bootstrap")
    fit = model_fit(res)
    fit_rows = []
    for metric in ("SRMR", "d_ULS", "d_G", "Chi_square", "NFI"):
        fit_rows.append([metric, fit["saturated"][metric], fit["estimated"][metric]])
    fit_rows.append(["RMS_theta", fit["RMS_theta"], np.nan])
    fit_rows.append(["GoF (Tenenhaus; descriptive only)", fit["GoF"], np.nan])

    def bollen_stine(kind):
        Sig = nearest_pd(fit["Sigma_sat" if kind == "saturated" else "Sigma_est"])
        Zbs = Zfull @ mat_power(fit["S"], -0.5) @ mat_power(Sig, 0.5)
        rng = np.random.default_rng(SEED + 7)
        dist = []
        for _ in range(N_BOOT_FIT):
            idx = rng.integers(0, n, n)
            Zb, _, _ = standardize(Zbs[idx])
            try:
                rb = pls_core(Zb, spec)
            except np.linalg.LinAlgError:
                continue
            Sb = (Zb.T @ Zb) / n
            Ib = implied_indicator_corr(rb.loadings, spec.ind_lv,
                                        rb.C if kind == "saturated" else implied_lv_corr(rb))
            dd = discrepancies(Sb, Ib, n)
            dist.append([dd["SRMR"], dd["d_ULS"], dd["d_G"]])
        return np.array(dist)

    exact = {}
    for kind in ("saturated", "estimated"):
        dist = bollen_stine(kind)
        exact[kind] = dist
        fig_fit_boot(dist[:, 0], fit[kind]["SRMR"], np.nanpercentile(dist[:, 0], 95),
                     np.nanpercentile(dist[:, 0], 99), fig_dir, figs, kind)
    fit_df = pd.DataFrame(fit_rows, columns=["Index", "Saturated_model", "Estimated_model"])
    for kind in ("saturated", "estimated"):
        d = exact[kind]
        hi95 = {"SRMR": np.nanpercentile(d[:, 0], 95), "d_ULS": np.nanpercentile(d[:, 1], 95),
                "d_G": np.nanpercentile(d[:, 2], 95)}
        hi99 = {"SRMR": np.nanpercentile(d[:, 0], 99), "d_ULS": np.nanpercentile(d[:, 1], 99),
                "d_G": np.nanpercentile(d[:, 2], 99)}
        col = "Saturated_model" if kind == "saturated" else "Estimated_model"
        fit_df[f"HI95_{kind}"] = fit_df["Index"].map(hi95)
        fit_df[f"HI99_{kind}"] = fit_df["Index"].map(hi99)
        fit_df[f"Exact_fit_95_{kind}"] = [
            (v <= hi95[m]) if m in hi95 else "" for m, v in zip(fit_df["Index"], fit_df[col])]
    fit_df["Guideline"] = fit_df["Index"].map({
        "SRMR": "< 0.08 (Hu & Bentler, 1999); < 0.10 lenient", "NFI": "> 0.90",
        "RMS_theta": "< 0.12", "d_ULS": "< HI95", "d_G": "< HI95",
        "Chi_square": "descriptive", "GoF (Tenenhaus; descriptive only)": "not a fit test"})
    add("T7_Model_fit", fit_df, f"Table 7. Model fit: SRMR, d_ULS, d_G, Chi², NFI (saturated and "
        f"estimated models), RMS_theta, and Bollen-Stine bootstrap exact-fit tests ({N_BOOT_FIT} runs).")
    summary.append(f"Model fit: SRMR saturated = {fit['saturated']['SRMR']:.3f}, estimated = "
                   f"{fit['estimated']['SRMR']:.3f}; NFI = {fit['saturated']['NFI']:.3f}; "
                   f"RMS_theta = {fit['RMS_theta']:.3f}.")

    # ---- CB-SEM with semopy (Table 8)
    cb_paths = None
    cb_fit = None
    try:
        log("CB-SEM with semopy: CFA, structural model, fit indices (Table 8)")
        cb = semopy_analysis(items, blocks, paths, out_dir)
        cb_fit = cb["Structural model"]["fit"]
        t8 = fit_indices_table(cb_fit, "structural model")
        add("T8_Structural_fit_indices", t8, "Table 8. Structural model fit indices (CB-SEM, maximum "
            "likelihood, semopy). *ATV = acceptable threshold value.")
        add("T8b_CFA_fit_indices", fit_indices_table(cb["Measurement model (CFA)"]["fit"], "measurement model"),
            "Fit indices of the measurement model (confirmatory factor analysis, semopy).")
        full = pd.DataFrame({"Index": list(cb_fit.keys()),
                             "Measurement_model_CFA": list(cb["Measurement model (CFA)"]["fit"].values()),
                             "Structural_model": list(cb_fit.values())})
        full["Guideline"] = full["Index"].map({
            "p-value": "> 0.05 (sensitive to N)", "X2/df": "1-5 (< 3 good)", "GFI": "> 0.90", "AGFI": "> 0.80",
            "PGFI": "> 0.50", "RMR": "small", "SRMR": "< 0.08", "NFI": "> 0.90", "RFI": "> 0.90",
            "IFI": "> 0.90", "TLI": "> 0.90", "CFI": "> 0.90 (> 0.95 good)", "PNFI": "> 0.50",
            "PCFI": "> 0.50", "RMSEA": "< 0.08 (< 0.10 acceptable)", "PCLOSE": "> 0.05",
            "AIC (semopy)": "lower is better", "BIC (semopy)": "lower is better"}).fillna("")
        add("T8c_CBSEM_all_fit_indices", full, "All CB-SEM fit indices for the measurement (CFA) and "
            "structural models (semopy ML estimation; GFI/AGFI/PGFI recomputed with the Joreskog-Sorbom formula).")
        add("CB1_CFA_estimates", cb["Measurement model (CFA)"]["estimates"],
            "CB-SEM confirmatory factor analysis: unstandardised and standardised estimates (semopy).")
        sem_est = cb["Structural model"]["estimates"]
        add("CB2_SEM_estimates", sem_est, "CB-SEM structural model: all parameter estimates (semopy).")
        lvs_ = list(blocks)
        reg = sem_est[(sem_est["op"] == "~") & sem_est["lval"].isin(lvs_) & sem_est["rval"].isin(lvs_)]
        cb_paths = {f"{rv} -> {lv}": (est, se, pv) for lv, rv, est, se, pv in
                    zip(reg["lval"], reg["rval"], reg["Est. Std"],
                        pd.to_numeric(reg["Std. Err"], errors="coerce") *
                        (pd.to_numeric(reg["Est. Std"]) / pd.to_numeric(reg["Estimate"])),
                        pd.to_numeric(reg["p-value"], errors="coerce"))}
        add("CB3_Estimator_robustness", cb["estimators"].reset_index().rename(columns={"index": "Path"}),
            "Standardised CB-SEM path coefficients under ML (MLW), ULS and GLS estimation (GLS is less "
            "stable with many indicators and moderate N; interpret ML as the reference).")
        cfa = cb["Measurement model (CFA)"]["estimates"]
        ld = cfa[(cfa["op"] == "~") & ~cfa["lval"].isin(list(blocks)) & cfa["rval"].isin(list(blocks))]
        lam_cb = dict(zip(ld["lval"], pd.to_numeric(ld["Est. Std"], errors="coerce")))
        cb_rel = []
        for c, v in blocks.items():
            lam = np.array([lam_cb[i] for i in v], dtype=float)
            cr = lam.sum() ** 2 / (lam.sum() ** 2 + (1 - lam ** 2).sum())
            cb_rel.append([c, cr, np.mean(lam ** 2), lam.min(), lam.max()])
        add("CB4_CFA_reliability", pd.DataFrame(cb_rel, columns=["Construct", "CR_(CFA)", "AVE_(CFA)",
                                                                  "Min_std_loading", "Max_std_loading"]),
            "Composite reliability and AVE from the CB-SEM CFA standardised loadings.")
        if cb["semplot"] is not None:
            figs.append(cb["semplot"])
        print("\nTable 8: Structural Model Fit Indices\n" + boxed_table(t8, header=True))
        summary.append("CB-SEM (semopy) structural model: " + ", ".join(
            f"{k} = {cb_fit[k]:.3f}" for k in ("X2/df", "GFI", "SRMR", "IFI", "NFI", "PGFI", "PNFI", "RMSEA", "CFI")))
    except ImportError:
        log("semopy is not installed - Table 8 skipped. Install with: pip install semopy")
    except Exception as exc:  # noqa: BLE001
        log(f"CB-SEM (semopy) step failed: {exc}")

    # ---- common method bias
    Cinv = np.linalg.pinv(res.C)
    fcvif = pd.DataFrame({"Construct": spec.lv, "Full_collinearity_VIF": np.diag(Cinv)})
    fcvif["<=3.3_(no_CMB)"] = fcvif["Full_collinearity_VIF"] <= 3.3
    add("C1_Full_collinearity_VIF", fcvif, "Full collinearity VIFs (Kock, 2015); VIF > 3.3 signals "
        "common method bias or collinearity.")
    summary.append(f"Common method bias: Harman first factor = {harman:.1f}% (concern if > 50%); "
                   f"max full-collinearity VIF = {fcvif['Full_collinearity_VIF'].max():.2f} (threshold 3.3).")

    # ---- measurement summary
    n_weak = int((load_df["Loading"] < 0.708).sum())
    bad_htmt = htmt_inf[~htmt_inf["HTMT<0.90"]]["Pair"].tolist()
    summary.insert(1, f"Measurement: {n_weak}/{k_i} loadings < 0.708; AVE < 0.50 for "
                      f"{rel.loc[~rel['AVE>=0.50'], 'Construct'].tolist() or 'none'}; "
                      f"Fornell-Larcker violated for {[c for c, ok in zip(spec.lv, fl_ok) if not ok] or 'none'}; "
                      f"HTMT > 0.90 for {bad_htmt or 'none'}.")
    summary.append("Hypotheses: " + "; ".join(
        f"{r.Hypothesis} {r.Path}: β={r.Original_sample:.3f}, t={r.t_statistic:.2f}, p {fmt_p(r.p_value)}"
        f" ({r.Decision})" for r in path_tab.itertuples()))

    # ---------------- robustness ----------------
    log("Robustness: Gaussian copula endogeneity test")
    cop, cop_norm = copula_test(res, N_BOOT_ROBUST, SEED)
    add("R1_Copula_endogeneity", cop, "Gaussian copula approach (Park & Gupta, 2012; Hult et al., 2018): "
        "significant copula terms indicate endogeneity.")
    add("R1b_Copula_identification", cop_norm, "Non-normality of predictors (required for copula identification).")
    log("Robustness: nonlinear effects")
    add("R2_Nonlinear_effects", nonlinear_test(res, N_BOOT_ROBUST, SEED),
        "Quadratic effects (Sarstedt et al., 2020): significant terms suggest nonlinearity.")

    # case influence (jackknife)
    jp = jack[:, sl["paths"]]
    infl = pd.DataFrame({"Case": np.arange(1, n + 1)})
    for k_, (s, t) in enumerate(spec.paths):
        infl[f"dBeta_{s}->{t}"] = est[k_] - jp[:, k_]
    infl["Max_abs_dBeta"] = infl.filter(like="dBeta").abs().max(axis=1)
    infl = infl.sort_values("Max_abs_dBeta", ascending=False)
    add("R3_Case_influence", infl, "Jackknife case influence: change in each path when a case is dropped.")

    # MGA
    for var, grp in MGA_GROUPS.items():
        if var not in demo.columns:
            continue
        codes = pd.to_numeric(demo[var], errors="coerce").values
        labels = list(grp.keys())
        g = np.full(n, -1)
        for k_, lab in enumerate(labels[:2]):
            g[np.isin(codes, grp[lab])] = k_
        keep = g >= 0
        sizes = [(g == 0).sum(), (g == 1).sum()]
        if min(sizes) < MGA_MIN_GROUP_SIZE:
            log(f"MGA by {var} skipped (group sizes {sizes} < {MGA_MIN_GROUP_SIZE})")
            continue
        log(f"Multigroup analysis by {var} (n = {sizes})")
        mt, mic = mga(X[keep], spec, g[keep], (labels[0], labels[1]), N_PERM, N_BOOT_ROBUST, SEED)
        add(f"R4_MGA_{var}", mt, f"PLS-MGA by {var} ({labels[0]} n={sizes[0]}, {labels[1]} n={sizes[1]}): "
            f"permutation test ({N_PERM} permutations) and group-specific bootstrap p-values.")
        add(f"R4b_MICOM_{var}", mic, f"MICOM steps 2-3 by {var} (Henseler et al., 2016).")
        fig_mga(mt, var, (labels[0], labels[1]), fig_dir, figs)

    # ---------------- sensitivity ----------------
    comp = path_tab[["Hypothesis", "Path"]].copy()
    comp["Main PLS_beta"] = path_tab["Original_sample"]
    comp["Main PLS_lo"] = path_tab["CI_2.5_percentile"]
    comp["Main PLS_hi"] = path_tab["CI_97.5_percentile"]
    comp["Main PLS_p"] = path_tab["p_value"]

    def path_extract(sp, ids):
        return lambda r: np.array([r.B[i, j] for i, j in ids])

    def run_variant(name, Xv, spv, ids=None, n_boot=N_BOOT_SENS):
        ids = ids or [(spv.lv_index[s], spv.lv_index[t]) for s, t in spec.paths]
        rv = fit_pls(Xv, spv)
        e = np.array([rv.B[i, j] for i, j in ids])
        bs = bootstrap(Xv, spv, n_boot, SEED + 1, path_extract(spv, ids), label=name)
        c = ci_table(e, bs, None, Xv.shape[0])
        comp[f"{name}_beta"] = e
        comp[f"{name}_lo"] = c["CI_2.5_percentile"].values
        comp[f"{name}_hi"] = c["CI_97.5_percentile"].values
        comp[f"{name}_p"] = c["p_value"].values
        return rv, c

    # PLSc
    log("Sensitivity: consistent PLS (PLSc)")
    pc = plsc_correct(res, R)

    def extract_plsc(r):
        Rb = (r.Z.T @ r.Z) / r.Z.shape[0]
        return np.array([plsc_correct(r, Rb)["B"][i, j] for i, j in path_ids])

    bs_c = bootstrap(X, spec, N_BOOT_SENS, SEED + 2, extract_plsc, label="PLSc")
    e_c = np.array([pc["B"][i, j] for i, j in path_ids])
    c_c = ci_table(e_c, bs_c, None, n)
    comp["PLSc_beta"], comp["PLSc_lo"], comp["PLSc_hi"], comp["PLSc_p"] = (
        e_c, c_c["CI_2.5_percentile"].values, c_c["CI_97.5_percentile"].values, c_c["p_value"].values)
    plsc_tab = pd.DataFrame({"Construct": spec.lv, "rho_A": pc["rho_A"],
                             "R2_PLSc": pc["R2"]})
    add("P1_PLSc_constructs", plsc_tab, "Consistent PLS: reliabilities used for disattenuation and R².")
    plsc_corr = pd.DataFrame(pc["C"], columns=spec.lv, index=spec.lv).reset_index().rename(
        columns={"index": "Construct"})
    add("P1b_PLSc_correlations", plsc_corr, "Disattenuated (PLSc) construct correlations; "
        "values >= 1 indicate a lack of discriminant validity.")

    # weighting schemes
    log("Sensitivity: inner weighting schemes")
    ws_rows = []
    for sch in ("path", "factorial", "centroid"):
        rr = fit_pls(X, spec, sch)
        ws_rows.append([sch] + [rr.B[i, j] for i, j in path_ids] + [rr.R2[j] for j in spec.endog])
    add("P2_Weighting_schemes", pd.DataFrame(ws_rows, columns=["Scheme"] + comp["Path"].tolist() +
                                             [f"R2_{spec.lv[j]}" for j in spec.endog]),
        "Path coefficients under path, factorial and centroid inner weighting.")

    # bootstrap seed stability
    log("Sensitivity: bootstrap seed stability")
    seed_rows = []
    for sd in (11, 222, 3333):
        bs = bootstrap(X, spec, N_BOOT_SENS // 2, sd, path_extract(spec, path_ids), label=f"seed {sd}")
        c = ci_table(est[sl["paths"]], bs, None, n)
        seed_rows.append([sd] + c["p_value"].tolist())
    add("P3_Bootstrap_seed_stability", pd.DataFrame(seed_rows, columns=["Seed"] + comp["Path"].tolist()),
        "p-values of paths with different bootstrap random seeds.")

    # sum-score OLS
    log("Sensitivity: sum-score regression (HC3)")
    Ms, _, _ = standardize(means.values)
    ss_rows = []
    for (s, t) in spec.paths:
        j = list(means.columns).index(t)
        P = [list(means.columns).index(a) for a, b in spec.paths if b == t]
        b, se, _ = ols(Ms[:, j], Ms[:, P])
        k_ = P.index(list(means.columns).index(s)) + 1
        tval = b[k_] / se[k_]
        ss_rows.append([b[k_], b[k_] - 1.96 * se[k_], b[k_] + 1.96 * se[k_], p_from_t(tval, n - len(P) - 1)])
    ss = np.array(ss_rows)
    comp["Sum-score OLS_beta"], comp["Sum-score OLS_lo"], comp["Sum-score OLS_hi"], comp["Sum-score OLS_p"] = ss.T

    if cb_paths:
        e_cb = np.array([cb_paths.get(pth, (np.nan,) * 3)[0] for pth in comp["Path"]], dtype=float)
        se_cb = np.array([cb_paths.get(pth, (np.nan,) * 3)[1] for pth in comp["Path"]], dtype=float)
        comp["CB-SEM (ML)_beta"] = e_cb
        comp["CB-SEM (ML)_lo"] = e_cb - 1.96 * se_cb
        comp["CB-SEM (ML)_hi"] = e_cb + 1.96 * se_cb
        comp["CB-SEM (ML)_p"] = [cb_paths.get(pth, (np.nan,) * 3)[2] for pth in comp["Path"]]

    # outliers removed
    if 0 < n_out < n * 0.2:
        log(f"Sensitivity: excluding {n_out} multivariate outliers")
        keep = ~md["Outlier_p<0.001"].values
        run_variant("Outliers removed", X[keep], spec)

    # control variables
    ctrl = [c for c in DEMOGRAPHICS if c in demo.columns and pd.to_numeric(demo[c], errors="coerce").notna().all()]
    if ctrl:
        log(f"Sensitivity: control variables {ctrl}")
        cblocks = OrderedDict(list(blocks.items()) + [(f"ctrl_{c}", [f"ctrl_{c}"]) for c in ctrl])
        cpaths = paths + [(f"ctrl_{c}", t) for c in ctrl for t in [spec.lv[j] for j in spec.endog]]
        cspec = ModelSpec(cblocks, cpaths)
        Xc = np.column_stack([X] + [pd.to_numeric(demo[c]).values for c in ctrl])
        rc, cc = run_variant("With controls", Xc, cspec)
        ctab = []
        for c in ctrl:
            for t in [spec.lv[j] for j in spec.endog]:
                ctab.append([c, t, rc.B[cspec.lv_index[f"ctrl_{c}"], cspec.lv_index[t]]])
        ctab = pd.DataFrame(ctab, columns=["Control", "Target", "beta"])
        ext = path_extract(cspec, [(cspec.lv_index[f"ctrl_{c}"], cspec.lv_index[t]) for c, t, _ in ctab.values])
        bs = bootstrap(Xc, cspec, N_BOOT_SENS, SEED + 3, ext, label="controls")
        cci = ci_table(ctab["beta"].values, bs, None, n)
        ctab["t"], ctab["p_value"] = cci["t_statistic"].values, cci["p_value"].values
        add("P4_Control_variables", ctab, "Effects of demographic control variables on accountability dimensions.")
        add("P4b_R2_with_controls", pd.DataFrame({"Construct": [spec.lv[j] for j in spec.endog],
                                                  "R2_main": [res.R2[j] for j in spec.endog],
                                                  "R2_with_controls": [rc.R2[cspec.lv_index[spec.lv[j]]]
                                                                       for j in spec.endog]}),
            "R² with and without demographic controls.")

    # mediation model CA -> FM -> accountability
    log("Alternative model: CA -> FM -> accountability (mediation)")
    mspec = ModelSpec(blocks, [("CA", "FM")] + paths)
    rm = fit_pls(X, mspec)
    ca, fm = mspec.lv_index["CA"], mspec.lv_index["FM"]
    targets = [mspec.lv_index[c] for c in ("ORG", "LEG", "PRO", "POL")]

    def ext_med(r):
        T = r.total
        return np.array([r.B[ca, fm]] + [r.B[ca, t] for t in targets] + [r.B[fm, t] for t in targets] +
                        [r.B[ca, fm] * r.B[fm, t] for t in targets] + [T[ca, t] for t in targets])

    e_m = ext_med(rm)
    bs_m = bootstrap(X, mspec, N_BOOT_SENS, SEED + 4, ext_med, label="mediation")
    cm = ci_table(e_m, bs_m, None, n)
    lab = (["CA -> FM"] + [f"CA -> {mspec.lv[t]} (direct)" for t in targets] +
           [f"FM -> {mspec.lv[t]}" for t in targets] +
           [f"CA -> FM -> {mspec.lv[t]} (indirect)" for t in targets] +
           [f"CA -> {mspec.lv[t]} (total)" for t in targets])
    med = pd.concat([pd.DataFrame({"Effect": lab}), cm], axis=1)
    vaf = []
    for k_, t in enumerate(targets):
        ind_v, tot_v = e_m[9 + k_], e_m[13 + k_]
        vaf.append([mspec.lv[t], ind_v, tot_v, ind_v / tot_v if tot_v != 0 else np.nan])
    add("A1_Mediation_model", med, "Alternative model with CA -> FM: direct, indirect and total effects "
        f"({N_BOOT_SENS} bootstrap subsamples).")
    add("A1b_Mediation_VAF", pd.DataFrame(vaf, columns=["Target", "Indirect", "Total", "VAF"]),
        "Variance accounted for (VAF = indirect / total).")
    mfit = model_fit(rm)
    add("A1c_Mediation_R2_fit", pd.DataFrame({
        "Construct": [mspec.lv[j] for j in mspec.endog], "R2": [rm.R2[j] for j in mspec.endog]}).assign(
        SRMR_estimated=mfit["estimated"]["SRMR"]), "R² and SRMR of the mediation model.")
    fig_structural(mspec, med.iloc[:9].assign(
        Source=["CA"] * 5 + ["FM"] * 4,
        Target=["FM"] + [mspec.lv[t] for t in targets] * 2, Hypothesis=""),
        {mspec.lv[j]: rm.R2[j] for j in mspec.endog}, fig_dir, figs,
        name="Fig19_mediation_model", title="Alternative model: CA → FM → accountability")

    # higher-order models (two-stage disjoint)
    log("Alternative models: higher-order constructs (two-stage)")
    lvs = pd.DataFrame(res.Y, columns=spec.lv)
    X2 = np.column_stack([lvs[["CA", "FM"]].values] + [items[blocks[c]].values for c in ("ORG", "LEG", "PRO", "POL")])
    hspec = ModelSpec(OrderedDict([("FC", ["CA_score", "FM_score"])] +
                                  [(c, blocks[c]) for c in ("ORG", "LEG", "PRO", "POL")]),
                      [("FC", c) for c in ("ORG", "LEG", "PRO", "POL")])
    rh = fit_pls(X2, hspec)
    ids_h = [(0, hspec.lv_index[c]) for c in ("ORG", "LEG", "PRO", "POL")]
    bs_h = bootstrap(X2, hspec, N_BOOT_SENS, SEED + 5, path_extract(hspec, ids_h), label="HOC FC")
    ch = ci_table(np.array([rh.B[i, j] for i, j in ids_h]), bs_h, None, n)
    hoc1 = pd.concat([pd.DataFrame({"Path": [f"FC -> {c}" for c in ("ORG", "LEG", "PRO", "POL")],
                                    "R2": [rh.R2[j] for _, j in ids_h]}), ch], axis=1)
    rel_h = reliability(rh, corr(X2))
    add("A2_HOC_financial_control", hoc1, "Higher-order model 1: Financial Control (CA + FM, two-stage "
        "disjoint approach) -> four accountability dimensions.")
    add("A2b_HOC_FC_measurement", pd.concat([
        pd.DataFrame({"Indicator": hspec.indicators[:2], "Loading": rh.loadings[:2]}),
        rel_h.iloc[[0]].reset_index(drop=True)], axis=1), "Measurement quality of the higher-order "
        "Financial Control construct.")
    X3 = np.column_stack([items[blocks["CA"]].values, items[blocks["FM"]].values,
                          lvs[["ORG", "LEG", "PRO", "POL"]].values])
    hspec2 = ModelSpec(OrderedDict([("CA", blocks["CA"]), ("FM", blocks["FM"]),
                                    ("ACC", ["ORG_score", "LEG_score", "PRO_score", "POL_score"])]),
                       [("CA", "ACC"), ("FM", "ACC")])
    rh2 = fit_pls(X3, hspec2)
    ids_h2 = [(0, 2), (1, 2)]
    bs_h2 = bootstrap(X3, hspec2, N_BOOT_SENS, SEED + 6, path_extract(hspec2, ids_h2), label="HOC ACC")
    ch2 = ci_table(np.array([rh2.B[i, j] for i, j in ids_h2]), bs_h2, None, n)
    rel_h2 = reliability(rh2, corr(X3))
    add("A3_HOC_accountability", pd.concat([pd.DataFrame({"Path": ["CA -> ACC", "FM -> ACC"],
                                                          "R2_ACC": rh2.R2[2]}), ch2], axis=1),
        "Higher-order model 2: CA and FM -> overall Public Accountability (second-order, two-stage).")
    add("A3b_HOC_ACC_measurement", pd.concat([
        pd.DataFrame({"Indicator": hspec2.indicators[-4:], "Loading": rh2.loadings[-4:]}),
        rel_h2.iloc[[2]].reset_index(drop=True)], axis=1),
        "Measurement quality of the higher-order Accountability construct.")

    # purified model
    log("Sensitivity: item purification (loadings + HTMT)")
    pblocks, plog = purify(items_all, ModelSpec(blocks_full, paths))
    add("P5_Purification_log", plog if len(plog) else pd.DataFrame([["No item removed"]], columns=["Step"]),
        f"Data-driven purification: weak loadings, then HTMT > {PURIFY_HTMT_TARGET} (min {PURIFY_MIN_ITEMS} "
        f"items, max {int(PURIFY_MAX_SHARE * 100)}% removed per construct). EXPLORATORY - justify every "
        "removal on content grounds before using it.")
    pspec = ModelSpec(pblocks, paths)
    Xp = items_all[pspec.indicators].values
    rp, _ = run_variant("Purified", Xp, pspec, n_boot=N_BOOT_SENS)
    Rp = corr(Xp)
    relp = reliability(rp, Rp)
    Hp = htmt_matrix(Rp, pspec)
    flp = []
    for j in range(pspec.L):
        flp.append(relp["sqrt_AVE"].iloc[j] > np.abs(np.delete(rp.C[j], j)).max())
    relp["Fornell_Larcker_met"] = flp
    relp["Max_HTMT"] = [np.nanmax(np.concatenate([Hp[j, :j], Hp[j + 1:, j]])) for j in range(pspec.L)]
    add("P5b_Purified_reliability", relp, "Purified model: reliability, convergent and discriminant validity.")
    add("P5c_Purified_HTMT", pd.DataFrame(Hp, columns=pspec.lv, index=pspec.lv).reset_index()
        .rename(columns={"index": "Construct"}), "Purified model: HTMT matrix.")
    add("P5d_Purified_loadings", pd.DataFrame({"Construct": [pspec.lv[j] for j in pspec.ind_lv],
                                               "Indicator": pspec.indicators, "Loading": rp.loadings}),
        "Purified model: outer loadings.")
    fpur = model_fit(rp)
    add("P5e_Purified_fit_R2", pd.DataFrame({
        "Construct": [pspec.lv[j] for j in pspec.endog], "R2": [rp.R2[j] for j in pspec.endog]}).assign(
        SRMR_saturated=fpur["saturated"]["SRMR"], SRMR_estimated=fpur["estimated"]["SRMR"]),
        "Purified model: R² and SRMR.")
    summary.append(f"Purified model removed {len(plog)} item(s): {plog['Item_removed'].tolist() if len(plog) else []}; "
                   f"max HTMT {np.nanmax(H):.3f} -> {np.nanmax(Hp):.3f}.")

    add("P6_Robustness_comparison", comp, "Path coefficients, 95% CIs and p-values across all model "
        "specifications.")
    fig_forest(comp, fig_dir, figs)

    # ---------------- IPMA ----------------
    log("Importance-performance map analysis")
    rng_ = SCALE_MAX - SCALE_MIN
    Xr = (X - SCALE_MIN) / rng_ * 100
    sdx = X.std(axis=0, ddof=0)
    wu = res.weights / sdx
    LVr = np.zeros((n, L))
    wr_all = np.zeros(k_i)
    for j, b in enumerate(spec.block_idx):
        wr = wu[b] / wu[b].sum()
        wr_all[b] = wr
        LVr[:, j] = Xr[:, b] @ wr
    Bu = np.zeros((L, L))
    for j in spec.endog:
        P = spec.pred[j]
        bb = np.linalg.lstsq(np.column_stack([np.ones(n), LVr[:, P]]), LVr[:, j], rcond=None)[0]
        Bu[P, j] = bb[1:]
    Tu = Bu @ np.linalg.inv(np.eye(L) - Bu)
    ipma_all = []
    for j in spec.endog:
        rows = [[spec.lv[i], Tu[i, j], LVr[:, i].mean()] for i in range(L) if Tu[i, j] != 0]
        dfc = pd.DataFrame(rows, columns=["Predictor", "Importance", "Performance"])
        dfc.insert(0, "Target", spec.lv[j])
        ipma_all.append(dfc)
        fig_ipma(dfc, spec.lv[j], fig_dir, figs)
        rows_i = []
        for i in range(L):
            if Tu[i, j] != 0:
                for idx in spec.block_idx[i]:
                    rows_i.append([spec.indicators[idx], wr_all[idx] * Tu[i, j], Xr[:, idx].mean()])
        dfi = pd.DataFrame(rows_i, columns=["Predictor", "Importance", "Performance"])
        dfi.insert(0, "Target", spec.lv[j])
        ipma_all.append(dfi.assign(Level="indicator"))
        fig_ipma(dfi, spec.lv[j], fig_dir, figs, level="indicator")
    add("I1_IPMA", pd.concat(ipma_all).fillna({"Level": "construct"}),
        "IPMA: importance (unstandardised total effects) and performance (0-100) at construct and indicator level.")

    # ---------------- figures of main model ----------------
    fig_structural(spec, path_tab, {spec.lv[j]: res.R2[j] for j in spec.endog}, fig_dir, figs)
    fig_full_model(spec, res, path_tab, fig_dir, figs)

    log("3D figures")
    fig3d_response_surfaces(res, fig_dir, figs)
    fig3d_path_bars(path_tab, fig_dir, figs)
    fig3d_htmt(H, spec.lv, fig_dir, figs)
    fig3d_scatter(res, fig_dir, figs)
    fig3d_cross_loadings(cl, spec, fig_dir, figs)
    if interactive_3d(res, path_tab, H, out_dir / "Interactive_3D_figures.html"):
        log("Interactive 3D HTML written")
    else:
        log("plotly not installed - interactive 3D HTML skipped (pip install plotly)")

    # ---------------- overview and writing ----------------
    checks = pd.DataFrame([
        ["Indicator loadings >= 0.708", f"{k_i - n_weak}/{k_i}", n_weak == 0],
        ["Cronbach's alpha >= 0.70", f"{int(rel['alpha>=0.70'].sum())}/{L}", bool(rel['alpha>=0.70'].all())],
        ["rho_A >= 0.70", f"{int(rel['rho_A>=0.70'].sum())}/{L}", bool(rel['rho_A>=0.70'].all())],
        ["rho_C within 0.70-0.95", f"{int(rel['rho_C_0.70-0.95'].sum())}/{L}", bool(rel['rho_C_0.70-0.95'].all())],
        ["AVE >= 0.50", f"{int(rel['AVE>=0.50'].sum())}/{L}", bool(rel['AVE>=0.50'].all())],
        ["Fornell-Larcker met", f"{sum(fl_ok)}/{L}", all(fl_ok)],
        ["HTMT < 0.85", f"{int(htmt_inf['HTMT<0.85'].sum())}/{len(htmt_inf)}", bool(htmt_inf['HTMT<0.85'].all())],
        ["HTMT < 0.90", f"{int(htmt_inf['HTMT<0.90'].sum())}/{len(htmt_inf)}", bool(htmt_inf['HTMT<0.90'].all())],
        ["HTMT upper CI < 1", f"{int(htmt_inf['Upper_CI<1'].sum())}/{len(htmt_inf)}", bool(htmt_inf['Upper_CI<1'].all())],
        ["Cross-loading gap > 0.10", f"{int((~cl_df['Problem_(own<=other+0.10)']).sum())}/{k_i}",
         bool((~cl_df['Problem_(own<=other+0.10)']).all())],
        ["Outer VIF < 5", f"{int((load_df['Outer_VIF'] < 5).sum())}/{k_i}", bool((load_df['Outer_VIF'] < 5).all())],
        ["Inner VIF < 3.3", f"{int((path_tab['Inner_VIF'] < 3.3).sum())}/{len(path_tab)}",
         bool((path_tab['Inner_VIF'] < 3.3).all())],
        ["Full collinearity VIF <= 3.3", f"{int(fcvif['<=3.3_(no_CMB)'].sum())}/{L}", bool(fcvif['<=3.3_(no_CMB)'].all())],
        ["Harman first factor < 50%", f"{harman:.1f}%", harman < 50],
        ["SRMR (saturated) < 0.08", f"{fit['saturated']['SRMR']:.3f}", fit['saturated']['SRMR'] < 0.08],
        ["SRMR (estimated) < 0.08", f"{fit['estimated']['SRMR']:.3f}", fit['estimated']['SRMR'] < 0.08],
        ["NFI (saturated) > 0.90", f"{fit['saturated']['NFI']:.3f}", fit['saturated']['NFI'] > 0.90],
        ["Q² > 0 for all endogenous", ", ".join(f"{q2[j]:.3f}" for j in spec.endog), all(q2[j] > 0 for j in spec.endog)],
        ["Hypotheses supported", f"{int((path_tab['Decision'] == 'Supported').sum())}/{len(path_tab)}",
         bool((path_tab['Decision'] == 'Supported').all())],
        *([["CB-SEM " + k + " (" + a + ")", f"{cb_fit[k]:.3f}", ok(cb_fit[k])]
           for k, a, ok in (("X2/df", "1-5", lambda v: 1 <= v <= 5), ("GFI", "> 0.9", lambda v: v > 0.9),
                            ("SRMR", "< 0.08", lambda v: v < 0.08), ("IFI", "> 0.9", lambda v: v > 0.9),
                            ("NFI", "> 0.9", lambda v: v > 0.9), ("PGFI", "> 0.5", lambda v: v > 0.5),
                            ("PNFI", "> 0.5", lambda v: v > 0.5), ("RMSEA", "< 0.1", lambda v: v < 0.1),
                            ("CFI", "> 0.9", lambda v: v > 0.9))] if cb_fit else []),
        ["Copula terms non-significant", f"{int((cop['p_value'] >= 0.05).sum())}/{len(cop)}",
         bool((cop['p_value'] >= 0.05).all())],
    ], columns=["Criterion", "Result", "Met"])
    tables = OrderedDict([("00_Checklist", (checks, "Overview of all quality criteria (Hair et al., 2019, 2022)."))]
                         + list(tables.items()))

    xlsx = out_dir / "PLS_SEM_All_Results.xlsx"
    write_excel(tables, xlsx)
    log(f"Excel workbook: {xlsx}")

    summary_path = out_dir / "Summary.txt"
    with open(summary_path, "w", encoding="utf-8") as fh:
        fh.write("PLS-SEM ANALYSIS SUMMARY\n" + "=" * 70 + "\n")
        fh.write(f"Input: {in_path}\nRun time: {time.time() - t_start:.0f} s\n\n")
        for line in summary:
            fh.write("- " + line + "\n")
        fh.write("\nQUALITY CHECKLIST\n")
        fh.write(checks.to_string(index=False) + "\n")
    doc_tables = [("Table 1. Demographic profile", demo_tab),
                  ("Table 3. Outer loadings", load_df[["Construct", "Indicator", "Loading", "Loading_t",
                                                      "Loading_p", "Outer_VIF"]]),
                  ("Table 4. Reliability and convergent validity", rel[["Construct", "Cronbach_alpha", "rho_A",
                                                                        "Composite_reliability_rho_C", "AVE"]]),
                  ("Table 5. HTMT", htmt_df), ("Table 6. Fornell-Larcker criterion", fl_df.drop(columns=["Fornell_Larcker_met"])),
                  ("Table 7. Model fit", fit_df[["Index", "Saturated_model", "Estimated_model", "Guideline"]]),
                  *([("Table 8. Structural model fit indices", tables["T8_Structural_fit_indices"][0])]
                    if "T8_Structural_fit_indices" in tables else []),
                  ("Table 9. Hypotheses testing", path_tab[["Hypothesis", "Path", "Original_sample", "Sample_mean",
                                                            "Std_dev", "t_statistic", "p_value", "CI_2.5_percentile",
                                                            "CI_97.5_percentile", "f2", "Decision"]]),
                  ("Table 10. R² and Q²", r2_tab[["Construct", "R2", "R2_adjusted", "Q2_blindfolding"]]),
                  ("Quality checklist", checks)]
    if write_docx(doc_tables, figs, summary, out_dir / "PLS_SEM_Report.docx"):
        log("Word report written")
    else:
        log("python-docx not installed - Word report skipped (pip install python-docx)")

    print("\n" + "=" * 78)
    print("SUMMARY")
    print("=" * 78)
    for line in summary:
        print("- " + line)
    print("\n" + checks.to_string(index=False))
    print(f"\nAll results saved in: {out_dir}  ({len(figs)} figures, {len(tables)} tables)")
    print(f"Total run time: {time.time() - t_start:.0f} s")


if __name__ == "__main__":
    main()
