# -*- coding: utf-8 -*-
import os, json, sys
import numpy as np, pandas as pd
from scipy import stats
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter as L
H = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, H)
import instrument as I
D = os.path.join(H, "..", "deliverables")
HDR = PatternFill("solid", fgColor="DDE6F0"); BOLD = Font(bold=True, name="Calibri")
def sheet(wb, name, rows, widths=None, header=True, rtl=True):
    ws = wb.create_sheet(name); ws.sheet_view.rightToLeft = rtl
    for r in rows: ws.append(r)
    if header:
        for c in ws[1]: c.font = BOLD; c.fill = HDR; c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, w in enumerate(widths or [], 1): ws.column_dimensions[L(i)].width = w
    return ws
def build(p):
    R = json.load(open(os.path.join(H, "..", f"paper{p}", "output", f"results_paper{p}.json"), encoding="utf-8"))
    df = pd.read_csv(os.path.join(H, "..", "data", f"paper{p}_data.csv"))
    cons = I.PAPER[p]; items = [c for k in cons for c, _ in I.ITEMS[k]]
    demo = [d for d in I.DEMO if d[0] in I.DEMO_PAPER[p] and d[1] in df.columns]
    cols = [d[1] for d in demo] + items
    df = df[cols].copy(); df.insert(0, "ID", range(1, len(df) + 1))
    wb = Workbook(); wb.remove(wb.active)
    n = len(df)
    sheet(wb, "راهنما", [["فایل داده — مقالهٔ %d" % p], ["⚠ این داده‌ها با شبیه‌سازی واقع‌نما تولید شده‌اند و هیچ پاسخ‌دهندهٔ واقعی ندارند."],
        [f"تعداد مشاهدات: {n}"], ["برگه‌ها: Data (دادهٔ خام) | Codebook (کدنامه) | Scoring (راهنمای نمره‌دهی) | Scores_formula (نمره‌ها با فرمول Excel) | Scores_values (نمره‌ها به‌صورت عدد) | Descriptives | Reliability | Paths (ضرایب مسیر)"],
        ["برای جایگزینی دادهٔ واقعی: همین ساختار ستونی را حفظ کنید و فایل را به‌صورت CSV در کد (DATA_PATH) بدهید."]], [120], header=False)
    ws = sheet(wb, "Data", [list(df.columns)] + df.values.tolist(), [8] + [11] * (len(df.columns) - 1), rtl=False); ws.freeze_panes = "B2"
    # کدنامه
    cb = [["کد متغیر", "برچسب فارسی", "سازه", "نوع", "کدگذاری / دامنه"]]
    cb.append(["ID", "شمارهٔ پاسخ‌دهنده", "—", "شناسه", "۱ تا %d" % n])
    for code, var, lab, typ, cod, _ in demo: cb.append([var, lab, "جمعیت‌شناختی", typ, cod])
    for k in cons:
        for c, t in I.ITEMS[k]: cb.append([c, t, I.CONSTRUCTS[k][0] + " (%s)" % k, "لیکرت ۵ درجه‌ای", "۱ = کاملاً مخالفم … ۵ = کاملاً موافقم"])
    sheet(wb, "Codebook", cb, [14, 80, 46, 18, 46])
    sc = [["سازه", "کد", "گویه‌ها", "فرمول نمره", "دامنه", "تفسیر"]]
    for k in cons:
        cs = [c for c, _ in I.ITEMS[k]]
        sc.append([I.CONSTRUCTS[k][0], k, "، ".join(cs), f"میانگین {len(cs)} گویه", "۱ تا ۵", "۱–۲٫۳۳ پایین | ۲٫۳۴–۳٫۶۷ متوسط | ۳٫۶۸–۵ بالا"])
    sc += [[], ["قواعد: گویهٔ معکوس ندارد. حداقل ۸۰٪ گویه‌های هر سازه باید پاسخ داده شده باشد، وگرنه نمرهٔ سازه گم‌شده است. نمرهٔ بالاتر یعنی شدت بیشتر سازه (در کیفیت تصمیم: کیفیت بهتر)."]]
    sheet(wb, "Scoring", sc, [44, 10, 40, 22, 10, 50])
    colidx = {c: i + 1 for i, c in enumerate(df.columns)}
    f_rows = [["ID"] + cons]; v_rows = [["ID"] + cons]
    for r in range(n):
        row = [r + 1]; vr = [r + 1]
        for k in cons:
            cs = [c for c, _ in I.ITEMS[k]]; a, b = L(colidx[cs[0]]), L(colidx[cs[-1]])
            row.append(f"=IF(COUNT(Data!{a}{r+2}:{b}{r+2})>={int(np.ceil(0.8*len(cs)))},AVERAGE(Data!{a}{r+2}:{b}{r+2}),\"\")")
            vr.append(round(float(df[cs].iloc[r].mean()), 4))
        f_rows.append(row); v_rows.append(vr)
    sheet(wb, "Scores_formula", f_rows, [8] + [12] * len(cons), rtl=False); sheet(wb, "Scores_values", v_rows, [8] + [12] * len(cons), rtl=False)
    ds = [["گویه", "میانگین", "انحراف معیار", "حداقل", "حداکثر", "چولگی", "کشیدگی"]]
    for c in items: x = df[c]; ds.append([c, round(x.mean(), 3), round(x.std(), 3), int(x.min()), int(x.max()), round(float(stats.skew(x)), 3), round(float(stats.kurtosis(x)), 3)])
    sheet(wb, "Descriptives", ds, [12] + [14] * 6, rtl=False)
    rl = [["سازه", "تعداد گویه", "آلفای کرونباخ", "پایایی ترکیبی (CR)", "AVE", "√AVE", "حداقل بار", "حداکثر بار"]]
    for k in cons: v = R["rel"][k]; rl.append([I.CONSTRUCTS[k][0], v["items"], round(v["alpha"], 3), round(v["cr"], 3), round(v["ave"], 3), round(v["sqrt_ave"], 3), round(v["lmin"], 3), round(v["lmax"], 3)])
    sheet(wb, "Reliability", rl, [48] + [16] * 7)
    pt = [["مسیر (وابسته ~ مستقل)", "b", "β", "خطای معیار", "z", "p"]]
    for k, v in R["paths"].items(): pt.append([k, *[None if (isinstance(v[a], float) and np.isnan(v[a])) else round(v[a], 4) for a in ("b", "beta", "se", "z", "p")]])
    sheet(wb, "Paths", pt, [26, 12, 12, 12, 12, 12], rtl=False)
    out = os.path.join(D, f"paper{p}", f"Paper{p}_Data.xlsx"); wb.save(out); print(out, df.shape)
for p in (1, 2): build(p)
